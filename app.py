import os, re, io, base64, hmac, hashlib, json, uuid, time, secrets, sqlite3, socket, ipaddress, tempfile, urllib.parse, csv
from datetime import datetime, timezone, timedelta

import requests
from bs4 import BeautifulSoup
from flask import Flask, render_template_string, request, jsonify, session, redirect, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix

# ===========================================================================
# SETUP
# Environment variables (set them on Render -> Environment, never in code):
#   SECRET_KEY         required: any long random text (keeps logins safe)
#   DATABASE_URL       recommended: Postgres link (Neon/Supabase) so accounts survive redeploys
#   GEMINI_API_KEY     AI provider 1 (also reads images, PDFs and voice)
#   GROQ_API_KEY       AI provider 2 (also does voice transcription with Whisper)
#   TAVILY_API_KEY     optional: best live web search
#   OPENROUTER_API_KEY + OPENROUTER_MODEL   optional: AI provider 3
#   GEMINI_MODEL / GROQ_MODEL / GROQ_WHISPER_MODEL   optional: force a model name (otherwise auto-detected)
#   PROVIDER_ORDER     optional, default "gemini,groq,openrouter"
#   UNSPLASH_ACCESS_KEY / PEXELS_API_KEY / PIXABAY_API_KEY   optional: free keys for real photo search
#
# Render start command (important, stops "worker timeout" errors on slow AI calls):
#   gunicorn app:app --workers 1 --threads 8 --timeout 120
# ===========================================================================
app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
SECRET = os.environ.get("SECRET_KEY")
if not SECRET:
    SECRET = secrets.token_hex(32)
    print("[warn] SECRET_KEY is not set. Logins will reset on every restart. Set it on Render.")
app.secret_key = SECRET
app.config.update(
    MAX_CONTENT_LENGTH=20 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(os.getenv("RENDER")),
    PERMANENT_SESSION_LIFETIME=timedelta(days=30),
)

GEMINI_KEY = (os.getenv("GEMINI_API_KEY") or "").strip()
GROQ_KEY = (os.getenv("GROQ_API_KEY") or "").strip()
TAVILY_KEY = (os.getenv("TAVILY_API_KEY") or "").strip()
OPENROUTER_KEY = (os.getenv("OPENROUTER_API_KEY") or "").strip()
OPENROUTER_MODEL = (os.getenv("OPENROUTER_MODEL") or "").strip()
UNSPLASH_KEY = (os.getenv("UNSPLASH_ACCESS_KEY") or "").strip()
PEXELS_KEY = (os.getenv("PEXELS_API_KEY") or "").strip()
PIXABAY_KEY = (os.getenv("PIXABAY_API_KEY") or "").strip()
DATABASE_URL = (os.getenv("DATABASE_URL") or "").strip()
SQLITE_PATH = os.getenv("SQLITE_PATH", "morr.db")
PROVIDER_ORDER = [p.strip() for p in os.getenv("PROVIDER_ORDER", "gemini,groq,openrouter").split(",") if p.strip()]
# Free and guest users use this order (default keeps Gemini, our best model, for Premium)
FREE_PROVIDER_ORDER = [p.strip() for p in os.getenv("FREE_PROVIDER_ORDER", "groq,gemini,openrouter").split(",") if p.strip()]


def _int_env(name, default):
    try:
        return int(os.getenv(name, default))
    except (TypeError, ValueError):
        return int(default)


# --- Monetisation settings (all changeable on Render without touching the code) ---
PAYSTACK_SECRET = (os.getenv("PAYSTACK_SECRET_KEY") or "").strip()
PAYSTACK_PLAN = (os.getenv("PAYSTACK_PLAN_CODE") or "").strip()
try:
    PRICE_GHS = float(os.getenv("PREMIUM_PRICE_GHS", "53"))
except ValueError:
    PRICE_GHS = 53.0
ADMIN_EMAILS = {e.strip().lower() for e in os.getenv("ADMIN_EMAILS", "").split(",") if e.strip()}
SITE_URL = (os.getenv("SITE_URL") or "").strip().rstrip("/")
SIGNUPS_PER_IP = _int_env("SIGNUPS_PER_IP", 3)
DAYS_ONE_TIME = 30       # Mobile Money / one-off payment
DAYS_SUBSCRIPTION = 32   # card subscription (a couple of spare days so a late renewal never cuts access)
LIMITS = {
    "guest": {"msg": _int_env("GUEST_MSGS", 5), "img": 0, "upload": 0, "voice": 0, "file": 0},
    "free": {"msg": _int_env("FREE_MSGS", 20), "img": _int_env("FREE_IMGS", 3),
             "upload": _int_env("FREE_UPLOADS", 3), "voice": _int_env("FREE_VOICE", 20),
             "file": _int_env("FREE_FILES", 3)},
    "premium": {"msg": _int_env("PREMIUM_MSGS", 500), "img": _int_env("PREMIUM_IMGS", 50),
                "upload": _int_env("PREMIUM_UPLOADS", 50), "voice": _int_env("PREMIUM_VOICE", 200),
                "file": _int_env("PREMIUM_FILES", 50)},
}
MAX_UPLOAD_MB = {"free": _int_env("FREE_UPLOAD_MB", 5), "premium": _int_env("PREMIUM_UPLOAD_MB", 20)}

gclient = None
types = None
if GEMINI_KEY:
    try:
        from google import genai
        from google.genai import types
        gclient = genai.Client(api_key=GEMINI_KEY)
    except Exception as e:
        print(f"[gemini] could not start client: {e}")

GENERATED_DIR = os.path.join(tempfile.gettempdir(), "morr_ai_generated")
os.makedirs(GENERATED_DIR, exist_ok=True)
BROWSER_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
WIKI_UA = {"User-Agent": "MORR AI GH/2.0 (Ghana AI assistant; image lookup)"}


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def cleanup_old_files(max_age=6 * 3600):
    try:
        cutoff = time.time() - max_age
        for name in os.listdir(GENERATED_DIR):
            p = os.path.join(GENERATED_DIR, name)
            if os.path.isfile(p) and os.path.getmtime(p) < cutoff:
                os.remove(p)
    except Exception:
        pass


# --- Cleans AI answers: removes leaked citation markers like 【1†L1-L4】 and odd invisible/narrow spaces ---
CITE_RE = re.compile(r"【[^】]*】")
_ODD_SPACES = dict.fromkeys(map(ord, "\u00a0\u1680\u2000\u2001\u2002\u2003\u2004\u2005\u2006\u2007\u2008\u2009\u200a\u202f\u205f\u3000"), " ")
_ZERO_WIDTH = dict.fromkeys(map(ord, "\u200b\u2060\ufeff"), None)


def clean_answer(text):
    if not text:
        return text
    text = CITE_RE.sub("", text)
    text = text.translate(_ODD_SPACES).translate(_ZERO_WIDTH)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"(?<=\S) +([.,])(?=\s|$)", r"\1", text)
    return text.strip()


# ===========================================================================
# DATABASE (Postgres if DATABASE_URL is set, otherwise SQLite)
# ===========================================================================
IS_PG = bool(DATABASE_URL)


def get_conn():
    if IS_PG:
        import psycopg2
        return psycopg2.connect(DATABASE_URL, connect_timeout=10)
    return sqlite3.connect(SQLITE_PATH)


def db(sql, params=(), fetch=None, insert=False):
    conn = get_conn()
    try:
        cur = conn.cursor()
        if IS_PG:
            sql = sql.replace("?", "%s")
            if insert:
                sql += " RETURNING id"
        cur.execute(sql, params)
        result = None
        if insert:
            result = cur.fetchone()[0] if IS_PG else cur.lastrowid
        elif fetch == "one":
            result = cur.fetchone()
        elif fetch == "all":
            result = cur.fetchall()
        conn.commit()
        return result
    finally:
        conn.close()


def init_db():
    pk = "SERIAL PRIMARY KEY" if IS_PG else "INTEGER PRIMARY KEY AUTOINCREMENT"
    try:
        db(f"CREATE TABLE IF NOT EXISTS users (id {pk}, email TEXT UNIQUE NOT NULL, password TEXT NOT NULL)")
        db(f"CREATE TABLE IF NOT EXISTS conversations (id {pk}, user_id INTEGER NOT NULL, title TEXT, created TEXT, updated TEXT)")
        db(f"CREATE TABLE IF NOT EXISTS messages (id {pk}, conversation_id INTEGER NOT NULL, role TEXT, content TEXT, created TEXT)")
        db("CREATE INDEX IF NOT EXISTS idx_conv_user ON conversations (user_id, updated)")
        db("CREATE INDEX IF NOT EXISTS idx_msg_conv ON messages (conversation_id)")
        db("CREATE TABLE IF NOT EXISTS subscriptions (user_id INTEGER PRIMARY KEY, premium_until TEXT, sub_code TEXT, "
           "email_token TEXT, auto_renew INTEGER DEFAULT 0, updated TEXT)")
        db("CREATE TABLE IF NOT EXISTS usage_daily (user_key TEXT NOT NULL, day TEXT NOT NULL, kind TEXT NOT NULL, "
           "n INTEGER NOT NULL DEFAULT 0, PRIMARY KEY (user_key, day, kind))")
        db(f"CREATE TABLE IF NOT EXISTS payments (id {pk}, reference TEXT UNIQUE NOT NULL, user_id INTEGER, email TEXT, "
           "amount INTEGER, kind TEXT, status TEXT, created TEXT)")
        if not IS_PG:
            print("[warn] DATABASE_URL not set: using a local SQLite file. Accounts and chats are LOST on every Render redeploy.")
    except Exception as e:
        print(f"[db] init failed: {e}")


init_db()


# ===========================================================================
# LIVE INTERNET SEARCH  (Tavily -> DuckDuckGo -> Wikipedia) + page reading
# ===========================================================================
def is_public_url(url):
    try:
        p = urllib.parse.urlparse(url)
        if p.scheme not in ("http", "https") or not p.hostname:
            return False
        for info in socket.getaddrinfo(p.hostname, None):
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
                return False
        return True
    except Exception:
        return False


def fetch_page(url, limit=1500):
    """Read the main text of a public web page (redirects are re-checked)."""
    try:
        for _ in range(4):
            if not is_public_url(url):
                return ""
            r = requests.get(url, headers={"User-Agent": BROWSER_UA}, timeout=6, allow_redirects=False)
            if r.status_code in (301, 302, 303, 307, 308):
                url = urllib.parse.urljoin(url, r.headers.get("Location", ""))
                continue
            if r.status_code != 200 or "text/html" not in r.headers.get("Content-Type", ""):
                return ""
            soup = BeautifulSoup(r.text[:300000], "html.parser")
            for t in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
                t.decompose()
            text = " ".join(p.get_text(" ", strip=True) for p in soup.find_all("p"))
            return re.sub(r"\s+", " ", text)[:limit]
    except Exception:
        pass
    return ""


def _ddgs_class():
    try:
        from ddgs import DDGS
        return DDGS
    except Exception:
        try:
            from duckduckgo_search import DDGS
            return DDGS
        except Exception:
            return None


def web_search(query, deep=False):
    """Returns {"text": str for the AI, "sources": [{"title","url"}]}"""
    items, direct = [], ""

    def add(title, snippet, url):
        if url and snippet and all(url != i["url"] for i in items):
            items.append({"title": (title or url)[:120], "snippet": snippet[:400], "url": url})

    if TAVILY_KEY:
        try:
            r = requests.post("https://api.tavily.com/search",
                              json={"api_key": TAVILY_KEY, "query": query, "max_results": 5,
                                    "include_answer": True, "search_depth": "basic"}, timeout=10)
            r.raise_for_status()
            data = r.json()
            direct = data.get("answer") or ""
            for x in data.get("results", []):
                add(x.get("title"), x.get("content"), x.get("url"))
        except Exception as e:
            print(f"[search] tavily failed: {str(e)[:200]}")

    if len(items) < 3:
        DDGS = _ddgs_class()
        if DDGS:
            try:
                for x in list(DDGS().text(query, max_results=6)):
                    add(x.get("title"), x.get("body"), x.get("href"))
            except Exception as e:
                print(f"[search] duckduckgo failed: {str(e)[:200]}")

    if len(items) < 3:
        try:
            r = requests.get("https://en.wikipedia.org/w/api.php", params={
                "action": "query", "format": "json", "generator": "search", "gsrsearch": query,
                "gsrlimit": 3, "prop": "extracts|info", "exintro": 1, "explaintext": 1,
                "exsentences": 4, "inprop": "url"}, headers={"User-Agent": "MORR AI GH/2.0"}, timeout=8)
            r.raise_for_status()
            pages = (r.json().get("query") or {}).get("pages", {})
            for p in sorted(pages.values(), key=lambda x: x.get("index", 99)):
                add(p.get("title"), p.get("extract"), p.get("fullurl"))
        except Exception as e:
            print(f"[search] wikipedia failed: {str(e)[:200]}")

    if deep:
        for it in items[:2]:
            it["page"] = fetch_page(it["url"])

    parts = []
    if direct:
        parts.append(f"Quick answer: {direct}")
    for n, it in enumerate(items[:6], 1):
        block = f"[{n}] {it['title']}\n{it['snippet']}"
        if it.get("page"):
            block += f"\nPage excerpt: {it['page']}"
        block += f"\nURL: {it['url']}"
        parts.append(block)
    return {"text": "\n\n".join(parts)[:7000],
            "sources": [{"title": i["title"], "url": i["url"]} for i in items[:4]]}


SMALLTALK = re.compile(
    r"^\s*(hi+|hello|hey+|yo|chale|charley|wagwan|how are you|how far|how you dey|how body|how is it going|"
    r"good (morning|afternoon|evening|night)|thanks?|thank you|ok(ay)?|cool|nice|bye|goodbye|who are you|"
    r"what('s| is) your name|what can you do|akwaaba|medaase|ete sen|sannu|ina kwana|na gode|bonjour|salut|hola|merci)\b",
    re.I)
NO_SEARCH = re.compile(r"^\s*(translate|rewrite|write|compose|draft|summari[sz]e|fix|debug|solve|calculate|"
                       r"correct|proofread|tell me a joke|joke)\b", re.I)
SEARCH_CUES = re.compile(
    r"\b(who is|who's|who was|president|prime minister|ceo|current|currently|latest|today|tonight|now|news|"
    r"price|weather|score|result|results|won|winner|when did|when is|where is|how much|population|capital|"
    r"exchange rate|release|stock|election|minister|governor|champion|2025|2026)\b", re.I)


def should_search(q, has_attachment):
    if not q or len(q.strip()) < 3:
        return False
    if SMALLTALK.match(q) or NO_SEARCH.match(q):
        return False
    if SEARCH_CUES.search(q):
        return True
    if has_attachment:
        return False
    words = q.split()
    return "?" in q or len(words) <= 6


# ===========================================================================
# AI PROVIDERS (models are auto-detected, so retired model names stop breaking the app)
# ===========================================================================
MODEL_CACHE = {}


COOLDOWN = {}  # model name -> time until which we skip it (after a 429 or 404)


def _cooling(model):
    return COOLDOWN.get(model, 0) > time.time()


def discover_gemini_models():
    """All usable Gemini flash models, newest first. 'lite' ones go last (they often have bigger free quotas)."""
    found = []
    try:
        for m in gclient.models.list():
            name = (m.name or "").replace("models/", "")
            if "gemini" not in name or "flash" not in name:
                continue
            if any(x in name for x in ("image", "tts", "live", "audio", "embed", "native", "robotics", "imagen")):
                continue
            actions = getattr(m, "supported_actions", None)
            if actions and "generateContent" not in actions:
                continue
            mt = re.search(r"gemini-(\d+(?:\.\d+)?)", name)
            ver = float(mt.group(1)) if mt else 0.0
            found.append(("lite" in name, -ver, len(name), name))
    except Exception as e:
        print(f"[gemini] model discovery failed: {str(e)[:200]}")
    found.sort()
    names = [f[3] for f in found]
    print(f"[gemini] usable models: {names[:6]}")
    return names


def _compose_models(discovered):
    lst = list(discovered)
    forced = os.getenv("GEMINI_MODEL")
    if forced:
        lst = [forced] + [m for m in lst if m != forced]
    return lst or ["gemini-3.8-flash"]


def gemini_models():
    lst = MODEL_CACHE.get("gemini_list")
    if lst is None:
        lst = _compose_models(discover_gemini_models() if gclient else [])
        MODEL_CACHE["gemini_list"] = lst
    return lst


def gemini_generate(contents, system=None, use_search=False):
    """Try every usable Gemini model in turn. Quota (429) and missing-model (404) errors move to the next one."""
    if not gclient:
        raise RuntimeError("Gemini is not configured")
    tried, last, refreshed = set(), None, False
    queue = list(gemini_models())
    while queue:
        model = queue.pop(0)
        if model in tried or _cooling(model):
            continue
        tried.add(model)
        try:
            cfg = types.GenerateContentConfig(
                system_instruction=system, temperature=0.7,
                tools=[types.Tool(google_search=types.GoogleSearch())] if use_search else None)
            return gclient.models.generate_content(model=model, contents=contents, config=cfg)
        except Exception as e:
            last = e
            s = str(e)
            if "NOT_FOUND" in s or "404" in s[:30]:
                COOLDOWN[model] = time.time() + 3600
                if not refreshed:
                    refreshed = True
                    fresh = discover_gemini_models() if gclient else []
                    if fresh:  # only replace the saved list if the fresh lookup actually worked
                        MODEL_CACHE["gemini_list"] = _compose_models(fresh)
                        queue.extend(m for m in MODEL_CACHE["gemini_list"] if m not in tried)
                continue
            if "429" in s[:30] or "RESOURCE_EXHAUSTED" in s:
                COOLDOWN[model] = time.time() + 120
                print(f"[gemini] {model} is over its quota, trying the next model")
                continue
            raise
    raise last or RuntimeError("All Gemini models are busy or unavailable right now")


def gemini_sources(resp):
    out = []
    try:
        for ch in (resp.candidates[0].grounding_metadata.grounding_chunks or [])[:4]:
            if ch.web and ch.web.uri:
                out.append({"title": ch.web.title or ch.web.uri, "url": ch.web.uri})
    except Exception:
        pass
    return out


def call_gemini(system, history, user_text):
    contents = [types.Content(role="user" if h["role"] == "user" else "model",
                              parts=[types.Part(text=h["content"])]) for h in history]
    contents.append(types.Content(role="user", parts=[types.Part(text=user_text)]))
    try:
        resp = gemini_generate(contents, system, use_search=True)
    except Exception as e:
        if "429" in str(e)[:30] or "RESOURCE_EXHAUSTED" in str(e):
            raise
        print(f"[gemini] search mode failed, retrying plain: {str(e)[:200]}")
        resp = gemini_generate(contents, system, use_search=False)
    if not resp.text:
        raise RuntimeError("Gemini returned an empty answer")
    return resp.text, gemini_sources(resp)


def _chat_post(url, key, model, messages, extra_headers=None, max_tokens=1200):
    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    headers.update(extra_headers or {})
    r = requests.post(url, headers=headers, timeout=45 if max_tokens <= 1500 else 90,
                      json={"model": model, "messages": messages, "temperature": 0.7, "max_tokens": max_tokens})
    if r.status_code != 200:
        raise RuntimeError(f"{r.status_code} {r.text[:300]}")
    return r.json()["choices"][0]["message"]["content"]


def discover_groq_model():
    try:
        r = requests.get("https://api.groq.com/openai/v1/models",
                         headers={"Authorization": f"Bearer {GROQ_KEY}"}, timeout=10)
        r.raise_for_status()
        best, best_size = None, -1
        for m in r.json().get("data", []):
            mid = m.get("id", "")
            if m.get("active") is False or any(x in mid.lower() for x in
                                               ("whisper", "guard", "tts", "playai", "embed", "orpheus", "vision-preview")):
                continue
            sizes = [float(x) for x in re.findall(r"(\d+(?:\.\d+)?)b", mid.lower())]
            size = max(sizes) if sizes else 0
            if size > best_size:
                best, best_size = mid, size
        print(f"[groq] auto-detected model: {best}")
        return best
    except Exception as e:
        print(f"[groq] model discovery failed: {str(e)[:200]}")
        return None


def call_groq(system, history, user_text, max_tokens=1200):
    messages = [{"role": "system", "content": system}] + history + [{"role": "user", "content": user_text}]
    forced = os.getenv("GROQ_MODEL")
    for attempt in range(2):
        model = MODEL_CACHE.get("groq") or forced or discover_groq_model()
        if not model:
            raise RuntimeError("No Groq model available")
        MODEL_CACHE["groq"] = model
        try:
            return _chat_post("https://api.groq.com/openai/v1/chat/completions", GROQ_KEY, model, messages,
                              max_tokens=max_tokens)
        except Exception as e:
            s = str(e).lower()
            if attempt == 0 and any(x in s for x in ("model_not_found", "decommissioned", "does not exist", "404")):
                MODEL_CACHE.pop("groq", None)
                forced = None
                continue
            raise


def call_openrouter(system, history, user_text, max_tokens=1200):
    messages = [{"role": "system", "content": system}] + history + [{"role": "user", "content": user_text}]
    return _chat_post("https://openrouter.ai/api/v1/chat/completions", OPENROUTER_KEY, OPENROUTER_MODEL,
                      messages, {"X-Title": "MORR AI GH"}, max_tokens)


# ===========================================================================
# VOICE TRANSCRIPTION (Groq Whisper first: it accepts webm/mp4/ogg recordings from phones and browsers)
# ===========================================================================
def discover_groq_whisper():
    try:
        r = requests.get("https://api.groq.com/openai/v1/models",
                         headers={"Authorization": f"Bearer {GROQ_KEY}"}, timeout=10)
        r.raise_for_status()
        ids = [m.get("id", "") for m in r.json().get("data", [])
               if "whisper" in m.get("id", "").lower() and m.get("active") is not False]
        for pref in ("whisper-large-v3-turbo", "whisper-large-v3"):
            if pref in ids:
                print(f"[groq] auto-detected whisper model: {pref}")
                return pref
        if ids:
            print(f"[groq] auto-detected whisper model: {ids[0]}")
            return ids[0]
    except Exception as e:
        print(f"[groq] whisper discovery failed: {str(e)[:200]}")
    return None


def groq_transcribe(data, filename, mime, lang_hint=None):
    if not GROQ_KEY:
        raise RuntimeError("Groq is not configured")
    forced = os.getenv("GROQ_WHISPER_MODEL")
    for attempt in range(2):
        model = MODEL_CACHE.get("groq_whisper") or forced or discover_groq_whisper() or "whisper-large-v3-turbo"
        MODEL_CACHE["groq_whisper"] = model
        form = {"model": model, "response_format": "json", "temperature": "0"}
        if lang_hint:
            form["language"] = lang_hint
        r = requests.post("https://api.groq.com/openai/v1/audio/transcriptions",
                          headers={"Authorization": f"Bearer {GROQ_KEY}"},
                          files={"file": (filename, data, mime)}, data=form, timeout=60)
        if r.status_code == 200:
            return (r.json().get("text") or "").strip()
        low = r.text[:300].lower()
        if attempt == 0 and any(x in low for x in ("model_not_found", "decommissioned", "does not exist")):
            MODEL_CACHE.pop("groq_whisper", None)
            forced = None
            continue
        raise RuntimeError(f"{r.status_code} {r.text[:200]}")


# ===========================================================================
# IMAGE READING (Gemini first, Groq vision as backup)
# ===========================================================================
def discover_groq_vision_model():
    try:
        r = requests.get("https://api.groq.com/openai/v1/models",
                         headers={"Authorization": f"Bearer {GROQ_KEY}"}, timeout=10)
        r.raise_for_status()
        ids = [m.get("id", "") for m in r.json().get("data", []) if m.get("active") is not False]
        for key in ("maverick", "scout", "vision"):
            for mid in ids:
                low = mid.lower()
                if key in low and not any(x in low for x in ("whisper", "guard", "tts")):
                    print(f"[groq] auto-detected vision model: {mid}")
                    return mid
    except Exception as e:
        print(f"[groq] vision model discovery failed: {str(e)[:200]}")
    return None


def groq_vision(jpeg_bytes, prompt):
    if not GROQ_KEY:
        raise RuntimeError("Groq is not configured")
    b64 = base64.b64encode(jpeg_bytes).decode()
    messages = [{"role": "user", "content": [
        {"type": "text", "text": prompt},
        {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + b64}}]}]
    forced = os.getenv("GROQ_VISION_MODEL")
    for attempt in range(2):
        model = MODEL_CACHE.get("groq_vision") or forced or discover_groq_vision_model()
        if not model:
            raise RuntimeError("No Groq vision model available")
        MODEL_CACHE["groq_vision"] = model
        try:
            return _chat_post("https://api.groq.com/openai/v1/chat/completions", GROQ_KEY, model, messages)
        except Exception as e:
            s = str(e).lower()
            if attempt == 0 and any(x in s for x in ("model_not_found", "decommissioned", "does not exist", "404")):
                MODEL_CACHE.pop("groq_vision", None)
                forced = None
                continue
            raise


def describe_image(jpeg_bytes):
    """Describe an image and read its text. Tries Gemini (all models), then Groq vision."""
    errors = []
    if gclient:
        try:
            resp = gemini_generate([types.Content(role="user", parts=[
                types.Part.from_bytes(data=jpeg_bytes, mime_type="image/jpeg"),
                types.Part(text=DESCRIBE_PROMPT)])])
            if resp.text:
                return resp.text
        except Exception as e:
            errors.append(f"gemini: {str(e)[:150]}")
            print(f"[vision] gemini failed: {str(e)[:200]}")
    if GROQ_KEY:
        try:
            text = groq_vision(jpeg_bytes, DESCRIBE_PROMPT)
            if text:
                return text
        except Exception as e:
            errors.append(f"groq: {str(e)[:150]}")
            print(f"[vision] groq failed: {str(e)[:200]}")
    raise RuntimeError("; ".join(errors) or "no image reader configured")


# ===========================================================================
# IMAGE REQUESTS: real photo search (Wikipedia + optional stock sites) or AI art (Pollinations / Gemini)
# ===========================================================================
IMG_GEN_RE = re.compile(
    r"\b(?:generate|create|make|draw|design|paint|produce|render)\b[^.?!\n]{0,40}?\b"
    r"(?:image|images|picture|pictures|photo|pic|logo|poster|illustration|wallpaper|drawing|painting|flyer|banner|icon|cartoon|artwork)\b"
    r"|\b(?:show me|give me|send me|find me|get me)\b[^.?!\n]{0,20}?\b(?:image|images|picture|pictures|photo|photos|pic|pics)\b"
    r"|\b(?:image|picture|photo|pic)s? of\b"
    r"|^\s*(?:/image|/imagine|imagine|draw)\b", re.I)
ART_WORDS = re.compile(r"\b(draw|paint|sketch|cartoon|anime|logo|poster|flyer|banner|icon|illustration|wallpaper|artwork|"
                       r"imagine|futuristic|fantasy|3d|abstract|watercolou?r|pixel art)\b", re.I)
PRONOUN_RE = re.compile(r"\b(his|her|hers|their|theirs|him|them|its|it)\b", re.I)
ATTACH_REF_RE = re.compile(
    r"\b(?:i\s+(?:just\s+)?(?:uploaded|attached|sent)|(?:uploaded|attached)\s+(?:picture|image|photo|file|pdf|document|screenshot)"
    r"|(?:this|the|my)\s+(?:picture|image|photo|screenshot|pdf|document|file)\s+(?:above|here))\b", re.I)


def wants_image(q):
    if re.match(r"^\s*(how|what|why|explain|where|when|who)\b", q, re.I):
        return False
    return bool(IMG_GEN_RE.search(q))


def clean_image_prompt(q):
    s = re.sub(r"^\s*(?:(?:please|pls|can you|could you)\s+)*"
               r"(?:(?:/image|/imagine|imagine|generate|create|make|draw|design|paint|produce|show me|give me|send me|find me|get me)\s+)?"
               r"(?:me\s+)?(?:an?\s+|the\s+)?(?:ai\s+)?(?:images?|pictures?|photos?|pics?|logo|poster|illustration|drawing|painting)?"
               r"\s*(?:of|for|showing|with)?\s*", "", q, flags=re.I).strip(" .:-")
    return s or q


def quick_llm(system, user):
    """One short AI call (Groq first, then Gemini). Returns '' if nothing is available."""
    if GROQ_KEY:
        try:
            return (call_groq(system, [], user) or "").strip()
        except Exception as e:
            print(f"[quick_llm] groq failed: {str(e)[:150]}")
    if gclient:
        try:
            resp = gemini_generate([types.Content(role="user", parts=[types.Part(text=user)])], system)
            return (resp.text or "").strip()
        except Exception as e:
            print(f"[quick_llm] gemini failed: {str(e)[:150]}")
    return ""


def plan_image(q, history):
    """Decide: show a REAL photo (searched online) or draw AI art. Uses the chat history so
    'his picture' becomes the real name. Returns {"mode": photo|art|clarify, "query": str, "person": bool}"""
    base = clean_image_prompt(q)
    ctx = "\n".join(f"{h['role']}: {h['content'][:300]}" for h in (history or [])[-4:])
    system = (
        "You decide how to fulfil an image request. Reply with ONLY one JSON object: "
        '{"mode":"photo" or "art","query":"...","person":true or false}. '
        'Use "photo" when the user wants a real, existing person, place, animal, object, event, product or landmark. '
        "Use the earlier chat to replace words like his/her/it/that with the real name. For photo, query is a short "
        "search phrase with the full name plus a disambiguating word, for example 'Ibrahim Traoré Burkina Faso' or "
        "'Independence Arch Accra'. Set person to true when the photo is of a real named human. "
        'Use "art" when they want something invented, artistic or hypothetical (drawings, logos, posters, cartoons, '
        "fantasy or imaginary scenes). For art, query is one vivid English image prompt of at most 50 words.")
    raw = quick_llm(system, f"Earlier chat:\n{ctx or '(none)'}\n\nImage request: {q}")
    m = re.search(r"\{.*\}", raw or "", re.S)
    if m:
        try:
            j = json.loads(m.group(0))
            mode = str(j.get("mode", "")).lower().strip()
            query = str(j.get("query", "")).strip().strip('"')
            if mode in ("photo", "art") and 2 < len(query) < 600:
                return {"mode": mode, "query": query, "person": bool(j.get("person"))}
        except Exception as e:
            print(f"[image] planner JSON failed: {str(e)[:120]}")
    # Fallback when no AI is reachable
    if ART_WORDS.search(q):
        return {"mode": "art", "query": base, "person": False}
    if PRONOUN_RE.search(base) and len(base.split()) <= 4:
        return {"mode": "clarify", "query": base, "person": False}
    return {"mode": "photo", "query": base, "person": False}


def _save_generated(data):
    from PIL import Image, ImageOps
    im = Image.open(io.BytesIO(data))
    im.load()
    im = ImageOps.exif_transpose(im).convert("RGB")
    im.thumbnail((1280, 1280))
    cleanup_old_files()
    name = f"gen_{uuid.uuid4().hex[:8]}.jpg"
    im.save(os.path.join(GENERATED_DIR, name), "JPEG", quality=90)
    return name


TRUSTED_IMG_HOSTS = ("wikimedia.org", "wikipedia.org", "images.unsplash.com", "images.pexels.com", "pixabay.com")


def _download_image(url, max_bytes=6 * 1024 * 1024):
    host = (urllib.parse.urlparse(url).hostname or "").lower()
    if not any(host == h or host.endswith("." + h) for h in TRUSTED_IMG_HOSTS):
        raise RuntimeError("untrusted image host")
    r = requests.get(url, headers=WIKI_UA, timeout=20, stream=True)
    r.raise_for_status()
    if not r.headers.get("Content-Type", "").startswith("image/"):
        raise RuntimeError("not an image")
    buf = b""
    for chunk in r.iter_content(65536):
        buf += chunk
        if len(buf) > max_bytes:
            raise RuntimeError("image too large")
    return buf


def wikipedia_photo(query):
    r = requests.get("https://en.wikipedia.org/w/api.php", params={
        "action": "query", "format": "json", "generator": "search", "gsrsearch": query, "gsrlimit": 5,
        "prop": "pageimages|info", "piprop": "thumbnail", "pithumbsize": 900, "pilimit": 5, "inprop": "url"},
        headers=WIKI_UA, timeout=10)
    r.raise_for_status()
    pages = (r.json().get("query") or {}).get("pages", {})
    for p in sorted(pages.values(), key=lambda x: x.get("index", 99)):
        th = (p.get("thumbnail") or {}).get("source")
        if th:
            return {"url": th, "credit": f"Wikipedia: {p.get('title')}", "page": p.get("fullurl")}
    return None


def unsplash_photo(query):
    if not UNSPLASH_KEY:
        return None
    r = requests.get("https://api.unsplash.com/search/photos", params={"query": query, "per_page": 1},
                     headers={"Authorization": f"Client-ID {UNSPLASH_KEY}"}, timeout=10)
    r.raise_for_status()
    res = r.json().get("results") or []
    if not res:
        return None
    x = res[0]
    return {"url": x["urls"]["regular"], "credit": f"Photo by {x['user']['name']} on Unsplash",
            "page": (x.get("links") or {}).get("html")}


def pexels_photo(query):
    if not PEXELS_KEY:
        return None
    r = requests.get("https://api.pexels.com/v1/search", params={"query": query, "per_page": 1},
                     headers={"Authorization": PEXELS_KEY}, timeout=10)
    r.raise_for_status()
    res = r.json().get("photos") or []
    if not res:
        return None
    x = res[0]
    return {"url": x["src"]["large"], "credit": f"Photo by {x.get('photographer', 'unknown')} on Pexels",
            "page": x.get("url")}


def pixabay_photo(query):
    if not PIXABAY_KEY:
        return None
    r = requests.get("https://pixabay.com/api/", params={"key": PIXABAY_KEY, "q": query, "image_type": "photo",
                                                         "per_page": 3, "safesearch": "true"}, timeout=10)
    r.raise_for_status()
    res = r.json().get("hits") or []
    if not res:
        return None
    x = res[0]
    return {"url": x["largeImageURL"], "credit": f"Image by {x.get('user', 'unknown')} on Pixabay",
            "page": x.get("pageURL")}


def find_real_photo(query, person=False):
    """Returns (saved_filename, credit, page_url, provider) or None. Real people: Wikipedia only (stock sites would
    return random strangers)."""
    order = [("wikipedia", wikipedia_photo)]
    if not person:
        order = [("unsplash", unsplash_photo), ("pexels", pexels_photo), ("pixabay", pixabay_photo)] + order
    for provider, fn in order:
        try:
            hit = fn(query)
            if not hit:
                continue
            data = _download_image(hit["url"])
            return _save_generated(data), hit["credit"], hit.get("page"), provider
        except Exception as e:
            print(f"[photo] {provider} failed: {str(e)[:150]}")
    return None


def pollinations_image(prompt):
    r = requests.get("https://image.pollinations.ai/prompt/" + urllib.parse.quote(prompt[:600]),
                     params={"width": 768, "height": 768, "nologo": "true", "seed": secrets.randbelow(10 ** 6)},
                     headers={"User-Agent": BROWSER_UA}, timeout=70)
    if r.status_code != 200 or not r.headers.get("Content-Type", "").startswith("image/"):
        raise RuntimeError(f"pollinations returned {r.status_code}")
    return r.content


def gemini_image(prompt):
    if not gclient:
        raise RuntimeError("Gemini is not configured")
    names = MODEL_CACHE.get("gemini_image")
    if names is None:
        names = []
        try:
            for m in gclient.models.list():
                n = (m.name or "").replace("models/", "")
                if "gemini" in n and "image" in n:
                    names.append(n)
        except Exception as e:
            print(f"[image] could not list Gemini image models: {str(e)[:150]}")
        MODEL_CACHE["gemini_image"] = names
    last = None
    for model in names[:2]:
        if _cooling(model):
            continue
        try:
            resp = gclient.models.generate_content(
                model=model, contents=prompt,
                config=types.GenerateContentConfig(response_modalities=["IMAGE", "TEXT"]))
            for part in resp.candidates[0].content.parts:
                inline = getattr(part, "inline_data", None)
                if inline and inline.data:
                    return inline.data
        except Exception as e:
            last = e
            COOLDOWN[model] = time.time() + 600
            print(f"[image] gemini {model} failed: {str(e)[:150]}")
    raise last or RuntimeError("Gemini returned no image")


def generate_image(prompt):
    errors = []
    for provider, fn in (("pollinations", pollinations_image), ("gemini", gemini_image)):
        try:
            return _save_generated(fn(prompt)), provider
        except Exception as e:
            errors.append(f"{provider}: {str(e)[:120]}")
            print(f"[image] {provider} failed: {str(e)[:200]}")
    raise RuntimeError("; ".join(errors))


def make_image_answer(q, info, art_ok=True):
    """Returns (answer, sources, provider, image_filename_or_None, usage_kind).
    Found photos count as a normal message; AI-drawn art counts against the image limit."""
    if info["mode"] == "clarify":
        return ("Who or what should the picture show? Please tell me the name, for example "
                "'photo of Ibrahim Traoré Burkina Faso'."), [], "image-clarify", None, "msg"
    if info["mode"] == "photo":
        try:
            res = find_real_photo(info["query"], info.get("person", False))
        except Exception as e:
            print(f"[photo] search crashed: {str(e)[:150]}")
            res = None
        if res:
            name, credit, page, prov = res
            ans = f"📷 Here is a real photo for: {info['query']}\nSource: {credit}" + (f" — {page}" if page else "")
            return ans, [], f"photo-{prov}", name, "msg"
        if info.get("person"):
            return (f"Sorry, I could not find a real photo of {info['query']}. I will not draw a made-up face of a real "
                    "person because it would not look like them. Try the full name and country, for example "
                    "'photo of Ibrahim Traoré Burkina Faso'."), [], "image-failed", None, "msg"
        if not art_ok:
            return ("I could not find a real photo of that, and you have used today's AI drawing limit. "
                    "Try different words or come back tomorrow."), [], "image-failed", None, "msg"
    if not art_ok:
        return "You have used today's AI drawing limit.", [], "image-failed", None, "img"
    prompt = info["query"]
    try:
        name, provider = generate_image(prompt)
        return f"🎨 Here is your AI-drawn image: {prompt[:100]}", [], f"image-{provider}", name, "img"
    except Exception:
        return ("Sorry, I could not draw that right now because the image services are busy. "
                "Please try again in a minute."), [], "image-failed", None, "img"


# ===========================================================================
# FILE GENERATION (the AI writes the content, then we build a real file you can download)
# ===========================================================================
FILE_GEN_RE = re.compile(
    r"\b(?:generate|create|make|write|produce|prepare|build|export|give me|draft|put (?:this|that|it) (?:in|into|as))\b[^.?!\n]{0,60}?\b"
    r"(pdf|word (?:document|doc|file)|docx|excel(?: file| sheet)?|xlsx|spreadsheet|csv|text file|txt file|txt)\b"
    r"|\b(?:generate|create|produce|export)\s+(?:an?\s+)?(?:\w+\s+){0,2}document\b", re.I)
FILE_LABEL = {"docx": "Word document", "pdf": "PDF", "xlsx": "Excel file", "csv": "CSV file", "txt": "text file"}
FILE_FALLBACK = {"pdf": ["docx", "txt"], "xlsx": ["csv", "txt"], "docx": ["txt"], "csv": ["txt"], "txt": []}


def wants_file(q):
    if re.match(r"^\s*(how|what|why|explain|where|when|who)\b", q, re.I):
        return False
    return bool(FILE_GEN_RE.search(q))


def detect_file_format(q):
    low = q.lower()
    if "pdf" in low:
        return "pdf"
    if re.search(r"excel|xlsx|spreadsheet", low):
        return "xlsx"
    if "csv" in low:
        return "csv"
    if re.search(r"\btxt\b|text file", low):
        return "txt"
    return "docx"


def plain(t):
    return re.sub(r"(\*\*|`|(?<!\w)\*(?!\s)|(?<!\s)\*(?!\w))", "", t or "").strip()


def md_blocks(md):
    """Very small Markdown reader: headings, bullets, tables, paragraphs."""
    out, para, table = [], [], []

    def flush_para():
        if para:
            out.append(("p", " ".join(para)))
            para.clear()

    def flush_table():
        if table:
            rows = [[c.strip() for c in r.strip().strip("|").split("|")]
                    for r in table if not re.match(r"^\s*\|?\s*:?-{2,}", r)]
            if rows:
                out.append(("table", rows))
            table.clear()

    for line in (md or "").splitlines():
        s = line.rstrip()
        if s.strip().startswith("|"):
            flush_para()
            table.append(s)
            continue
        flush_table()
        if not s.strip():
            flush_para()
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            flush_para()
            out.append((f"h{len(m.group(1))}", m.group(2).strip()))
            continue
        m = re.match(r"^\s*(?:[-*•]|\d+[.)])\s+(.*)", s)
        if m:
            flush_para()
            out.append(("li", m.group(1).strip()))
            continue
        para.append(s.strip())
    flush_para()
    flush_table()
    return out


def rows_from_md(md):
    blocks = md_blocks(md)
    for t, v in blocks:
        if t == "table":
            return v
    rows = [[v] for t, v in blocks if t != "table"]
    return rows or [[(md or "").strip()]]


def build_txt(md, path):
    lines = []
    for t, v in md_blocks(md):
        if t.startswith("h"):
            lines += [plain(v).upper() if t == "h1" else plain(v), ""]
        elif t == "li":
            lines.append("- " + plain(v))
        elif t == "table":
            lines += [" | ".join(plain(c) for c in r) for r in v] + [""]
        else:
            lines += [plain(v), ""]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines).strip() + "\n")


def build_csv(md, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        for r in rows_from_md(md):
            w.writerow([plain(c) for c in r])


def build_docx(md, path):
    import docx
    d = docx.Document()
    for t, v in md_blocks(md):
        if t == "h1":
            d.add_heading(plain(v), 0)
        elif t == "h2":
            d.add_heading(plain(v), 1)
        elif t == "h3":
            d.add_heading(plain(v), 2)
        elif t == "li":
            d.add_paragraph(plain(v), style="List Bullet")
        elif t == "table":
            n = max(len(r) for r in v)
            tbl = d.add_table(rows=len(v), cols=n)
            tbl.style = "Table Grid"
            for i, r in enumerate(v):
                for j in range(n):
                    tbl.cell(i, j).text = plain(r[j]) if j < len(r) else ""
            d.add_paragraph("")
        else:
            d.add_paragraph(plain(v))
    d.save(path)


def build_xlsx(md, path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in rows_from_md(md):
        ws.append([plain(c) for c in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    wb.save(path)


def build_pdf(md, path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    font, bold = "Helvetica", "Helvetica-Bold"
    try:  # DejaVu (if installed) shows letters like ɛ and ɔ used in Twi
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        reg = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        bld = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        if os.path.exists(reg) and os.path.exists(bld):
            pdfmetrics.registerFont(TTFont("MorrSans", reg))
            pdfmetrics.registerFont(TTFont("MorrSans-Bold", bld))
            font, bold = "MorrSans", "MorrSans-Bold"
    except Exception:
        pass
    st = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=st["BodyText"], fontName=font, fontSize=11, leading=15)
    h1 = ParagraphStyle("h1", parent=st["Title"], fontName=bold, fontSize=20, leading=24)
    h2 = ParagraphStyle("h2", parent=st["Heading2"], fontName=bold)
    h3 = ParagraphStyle("h3", parent=st["Heading3"], fontName=bold)
    li = ParagraphStyle("li", parent=body, leftIndent=16, bulletIndent=4)
    story = []
    for t, v in md_blocks(md):
        if t == "h1":
            story.append(Paragraph(escape(plain(v)), h1))
        elif t == "h2":
            story.append(Paragraph(escape(plain(v)), h2))
        elif t == "h3":
            story.append(Paragraph(escape(plain(v)), h3))
        elif t == "li":
            story.append(Paragraph(escape(plain(v)), li, bulletText="•"))
        elif t == "table":
            n = max(len(r) for r in v)
            data = [[Paragraph(escape(plain(c)), body) for c in (r + [""] * (n - len(r)))] for r in v]
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                     ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                                     ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [tbl, Spacer(1, 8)]
        else:
            story.append(Paragraph(escape(plain(v)), body))
        story.append(Spacer(1, 4))
    if not story:
        raise RuntimeError("empty document")
    SimpleDocTemplate(path, pagesize=A4, title="MORR AI GH document").build(story)


BUILDERS = {"pdf": build_pdf, "docx": build_docx, "xlsx": build_xlsx, "csv": build_csv, "txt": build_txt}


def build_file(md, fmt, title):
    """Returns (filename, format_used, note). Falls back to an easier format if a library is missing."""
    cleanup_old_files()
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", title).strip("_")[:40] or "document"
    last = None
    for f in [fmt] + FILE_FALLBACK.get(fmt, []):
        name = f"morr_{slug}_{uuid.uuid4().hex[:6]}.{f}"
        try:
            BUILDERS[f](md, os.path.join(GENERATED_DIR, name))
            note = "" if f == fmt else f" (I made a {FILE_LABEL[f]} instead because the {FILE_LABEL[fmt]} tools are not installed on the server)"
            return name, f, note
        except Exception as e:
            last = e
            print(f"[file] {f} build failed: {str(e)[:200]}")
    raise last or RuntimeError("could not build the file")


def file_system_prompt(lang, fmt):
    lang_rule = f"Write in {lang}." if lang and lang != "Auto" else "Write in the same language the user writes in."
    table_rule = ("The user wants a spreadsheet, so put the data in ONE Markdown table (header row, separator row, "
                  "then the rows) under a one-line '# Title'. " if fmt in ("xlsx", "csv") else "")
    return ("You write the full content of a file for the user. Output ONLY the file content in simple Markdown, "
            "with no commentary and no code fences. Start with '# Title', then '## Section' headings, short "
            "paragraphs and '- ' bullet lists. " + table_rule +
            "Use the web results and any attached content for facts. Never invent facts, names, numbers or links. " + lang_rule)


def generate_file_text(q, lang, file_ctx, img_ctx, history, order, fmt):
    topic = re.sub(r"\b(?:please|pls|can you|could you|generate|create|make|write|produce|prepare|build|export|draft|give me)\b",
                   " ", q, flags=re.I)
    topic = re.sub(r"\b(?:an?|the)?\s*(?:pdf|word (?:document|doc|file)|docx|excel(?: file| sheet)?|xlsx|spreadsheet|csv|text file|txt file|txt|document)\b",
                   " ", topic, flags=re.I)
    topic = re.sub(r"\s+", " ", topic).strip()
    web = {"text": ""}
    if len(topic) > 8 and not (file_ctx or img_ctx):
        try:
            web = web_search(topic)
        except Exception as e:
            print(f"[file] web search failed: {str(e)[:150]}")
    system = file_system_prompt(lang, fmt)
    user_text = build_user_text(q, file_ctx, img_ctx, web["text"])
    for provider in order:
        text = ""
        try:
            if provider == "gemini" and gclient:
                contents = [types.Content(role="user" if h["role"] == "user" else "model",
                                          parts=[types.Part(text=h["content"])]) for h in history]
                contents.append(types.Content(role="user", parts=[types.Part(text=user_text)]))
                text = gemini_generate(contents, system).text or ""
            elif provider == "groq" and GROQ_KEY:
                text = call_groq(system, history, user_text, max_tokens=3500)
            elif provider == "openrouter" and OPENROUTER_KEY and OPENROUTER_MODEL:
                text = call_openrouter(system, history, user_text, max_tokens=3500)
        except Exception as e:
            print(f"[file] {provider} failed: {str(e)[:200]}")
        text = clean_answer(text)
        text = re.sub(r"^```[a-zA-Z]*\s*\n|\n```\s*$", "", text.strip())
        if len(text.strip()) > 20:
            return text.strip()
    return ""


def make_file_answer(q, lang, file_ctx, img_ctx, history, order):
    """Returns (answer, sources, provider, filename_or_None)."""
    fmt = detect_file_format(q)
    md = generate_file_text(q, lang, file_ctx, img_ctx, history, order, fmt)
    if not md:
        return ("Sorry, I could not write that file right now because the AI services are busy. "
                "Please try again in a minute."), [], "file-failed", None
    m = re.search(r"^#\s+(.+)$", md, re.M)
    title = plain(m.group(1)) if m else "Your document"
    try:
        name, used, note = build_file(md, fmt, title)
    except Exception as e:
        print(f"[file] build failed: {str(e)[:200]}")
        return "Sorry, I wrote the content but could not build the file. Please try again.", [], "file-failed", None
    preview = re.sub(r"\s+", " ", " ".join(plain(v if isinstance(v, str) else " ".join(" ".join(r) for r in v))
                                          for t, v in md_blocks(md)))[:450]
    ans = f"📄 Your {FILE_LABEL[used]} is ready: {title}.{note}\n\nPreview: {preview}..."
    return ans, [], "file-ok", name


def friendly_name(filename):
    return re.sub(r"^morr_", "", re.sub(r"_[0-9a-f]{6,8}(?=\.[A-Za-z0-9]+$)", "", filename))


def build_system(lang):
    today = datetime.now(timezone.utc).strftime("%A, %d %B %Y")
    if lang and lang != "Auto":
        lang_rule = f"Reply in {lang}."
    else:
        lang_rule = ("Reply in the same language the user writes in (English, Ghanaian Pidgin, Twi, Hausa, French, "
                     "Arabic, Spanish or any other). If they mix languages, match their style.")
    return (
        "You are MORR AI GH, a smart, warm and friendly assistant built by Morr in Ghana, "
        "like ChatGPT or Meta AI. Talk like a real person: natural, direct, not robotic.\n"
        f"Today's date is {today}. {lang_rule}\n"
        "Rules:\n"
        "- Greetings and small talk get a short, human reply (for example 'how are you doing' -> "
        "'Chale, I'm good! How about you?'). Never search or define the greeting.\n"
        "- For facts, use the web results or search you are given and trust them over your memory. "
        "If the user's question has a wrong premise, politely correct it "
        "(for example: Jamaica has a Prime Minister, not a President).\n"
        "- If you are not sure, say so plainly. Never invent facts, names, numbers or links.\n"
        "- If a file or image is attached, use it first. Only say you cannot see an image or file if none was attached.\n"
        "- You can also find real photos, draw images and make PDF, Word, Excel, CSV and text files. If the user wants "
        "one, tell them to say, for example, 'photo of Accra Independence Arch', 'generate an image of a sunset in Accra' "
        "or 'create a PDF about healthy eating'.\n"
        "- Be concise. Use short paragraphs, and lists only when they help."
    )


def build_user_text(q, file_ctx, img_ctx, web_text=""):
    parts = []
    if web_text:
        parts.append("LIVE WEB SEARCH RESULTS (use these for facts, mention the source name when helpful):\n" + web_text)
    if file_ctx:
        parts.append("ATTACHED FILE CONTENT:\n" + file_ctx[:6000])
    if img_ctx:
        parts.append("ATTACHED IMAGE (description and any text in it):\n" + img_ctx[:3000])
    parts.append("USER MESSAGE:\n" + q)
    return "\n\n".join(parts)


def clean_history(hist):
    out = []
    for h in (hist or [])[-10:]:
        if isinstance(h, dict) and h.get("role") in ("user", "assistant") and isinstance(h.get("content"), str):
            out.append({"role": h["role"], "content": h["content"][:1500]})
    while out and out[0]["role"] != "user":
        out.pop(0)
    return out


def generate_answer(q, lang, file_ctx, img_ctx, history, order=None):
    """Try each AI provider in turn. Returns (answer, sources, provider)."""
    system = build_system(lang)
    has_att = bool(file_ctx or img_ctx)
    need_web = should_search(q, has_att)
    web = {"text": "", "sources": []}
    web_done = False

    for provider in (order or PROVIDER_ORDER):
        try:
            if provider == "gemini" and gclient:
                text, srcs = call_gemini(system, history, build_user_text(q, file_ctx, img_ctx))
                return clean_answer(text), srcs, "gemini"
            if provider in ("groq", "openrouter"):
                if provider == "groq" and not GROQ_KEY:
                    continue
                if provider == "openrouter" and not (OPENROUTER_KEY and OPENROUTER_MODEL):
                    continue
                if need_web and not web_done:
                    web = web_search(q, deep=True)
                    web_done = True
                user_text = build_user_text(q, file_ctx, img_ctx, web["text"])
                fn = call_groq if provider == "groq" else call_openrouter
                text = fn(system, history, user_text)
                if text:
                    return clean_answer(text), web["sources"], provider
        except Exception as e:
            print(f"[{provider}] failed: {str(e)[:300]}")

    # Last resort: show live search results directly, with a clear message
    if not web_done and q and not SMALLTALK.match(q):
        web = web_search(q)
    if web["text"]:
        top = web["text"].split("\n\n")[:3]
        msg = ("My AI brain is not reachable right now, but here is what I found on the web:\n\n"
               + "\n\n".join(top))
        return msg, web["sources"], "search-only"
    return ("Sorry, I could not reach my AI services or the internet just now. "
            "Please try again in a moment."), [], "none"


# ===========================================================================
# TEXT TO SPEECH
# ===========================================================================
def make_audio(text, lang):
    try:
        from gtts import gTTS
        plain_text = re.sub(r"https?://\S+|[*#_`>\[\]]", "", text).strip()[:350]
        if not plain_text or lang == "Twi":
            return None
        code = {"English": "en", "Pidgin": "en", "Hausa": "ha"}.get(lang)
        if code is None:
            code = "en" if re.fullmatch(r"[\x00-\x7F]+", plain_text, re.S) else None
        if code is None:
            return None
        cleanup_old_files()
        name = f"morr_{uuid.uuid4().hex[:8]}.mp3"
        gTTS(text=plain_text, lang=code).save(os.path.join(GENERATED_DIR, name))
        return f"/audio/{name}"
    except Exception as e:
        print(f"[tts] failed: {str(e)[:150]}")
        return None


# ===========================================================================
# PAGES
# ===========================================================================
HTML = r"""<!DOCTYPE html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>MORR AI GH</title>
<style>
*{box-sizing:border-box}
body{margin:0;display:flex;height:100vh;height:100dvh;font-family:system-ui,-apple-system,sans-serif;background:#0a0a0a;color:white}
#sidebar{width:272px;background:#171717;padding:14px;overflow-y:auto;border-right:1px solid #222;display:flex;flex-direction:column;flex-shrink:0}
#main{flex:1;display:flex;flex-direction:column;min-width:0}
#topbar{display:none;align-items:center;gap:10px;padding:10px 14px;border-bottom:1px solid #222}
#topbar b{color:#FFD700}
#chat{flex:1;overflow-y:auto;padding:16px;max-width:900px;width:100%;margin:0 auto}
.msg{padding:12px 16px;border-radius:18px;margin:10px 0;max-width:88%;line-height:1.6;white-space:pre-wrap;word-wrap:break-word;overflow-wrap:anywhere}
.user{background:#FFD700;color:black;margin-left:auto;border-bottom-right-radius:4px}
.ai{background:#1e1e1e;border:1px solid #333;border-bottom-left-radius:4px}
.msg img{max-width:100%;border-radius:12px;margin-top:8px}
.msg a{color:#FFD700}
.user a{color:#5a4600}
.msg code{background:#000;color:#fff;padding:1px 5px;border-radius:5px}
.src{display:block;font-size:12px;margin-top:8px;color:#aaa;white-space:normal}
.input-area{border-top:1px solid #222;padding:10px;background:#0a0a0a}
.row{display:flex;gap:6px;align-items:center;max-width:900px;margin:0 auto;width:100%}
select{background:#2a2a2a;color:white;padding:10px 6px;border-radius:10px;border:none;max-width:92px}
#q{flex:1;padding:13px;border-radius:12px;border:none;background:#1e1e1e;color:white;font-size:16px;min-width:0}
.btn{padding:11px 13px;background:#FFD700;color:black;border:none;border-radius:12px;font-weight:bold;cursor:pointer}
.btn2{background:#333;color:white}
.side-btn{background:#222;color:#ddd;border:1px solid #333;padding:10px;border-radius:10px;cursor:pointer;width:100%;font-size:14px;margin:8px 0}
.grp{font-size:11px;color:#777;margin:14px 0 4px;text-transform:uppercase}
.hist{display:flex;align-items:center;justify-content:space-between;padding:8px 10px;border-radius:8px;cursor:pointer;color:#bbb;font-size:13px;gap:6px}
.hist:hover,.hist.active{background:#262626;color:white}
.usage{margin:8px 0;padding:8px;background:#1f1f1f;border:1px solid #2c2c2c;border-radius:8px;font-size:12px;color:#bbb;line-height:1.6}.usage a{color:#FFD700}
.hist span{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.hist i{font-style:normal;color:#777;padding:0 4px}
.hist i:hover{color:#ff6b6b}
#overlay{display:none}
@media(max-width:800px){
 #sidebar{position:fixed;left:0;top:0;bottom:0;z-index:20;transform:translateX(-100%);transition:transform .2s}
 #sidebar.open{transform:none}
 #overlay.open{display:block;position:fixed;inset:0;background:#0009;z-index:10}
 #topbar{display:flex}
 .msg{max-width:94%}
}
</style></head><body>
<div id="overlay" onclick="toggleSide()"></div>
<div id="sidebar">
 <h3 style="color:#FFD700;margin:0 0 6px">MORR AI GH 🇬🇭</h3>
 {% if user_email %}
 <div style="margin:8px 0;padding:8px;background:#222;border-radius:8px;font-size:13px;word-break:break-all">{{user_email}}<br><a href="/logout" style="color:#FFD700">Logout</a></div>
 {% else %}
 <div style="margin:8px 0;font-size:12px;color:#999">Create a free account to save your chats and open them later.</div>
 <a href="/login" style="display:block;padding:10px;background:#FFD700;color:black;text-align:center;border-radius:8px;font-weight:bold;text-decoration:none;margin-bottom:6px">Log in</a>
 <a href="/signup" style="display:block;padding:10px;background:#222;color:white;text-align:center;border-radius:8px;text-decoration:none;border:1px solid #333">Create account</a>
 {% endif %}
 <div id="usageBox" class="usage"></div>
 <button class="side-btn" onclick="newChat()">＋ New chat</button>
 <div id="histList"></div>
 <div style="margin-top:auto;padding-top:14px;font-size:11px;color:#666"><a href="/pricing" style="color:#999">Plans</a> · <a href="/terms" style="color:#999">Terms</a> · <a href="/privacy" style="color:#999">Privacy</a></div>
</div>
<div id="main">
 <div id="topbar"><button class="btn btn2" onclick="toggleSide()">☰</button><b>MORR AI GH</b></div>
 <div id="chat"></div>
 <div class="input-area">
  <div id="status" style="color:#FFD700;font-size:12px;max-width:900px;margin:0 auto 6px;min-height:14px"></div>
  <div class="row">
   <select id="lang"><option value="Auto">Auto</option><option>English</option><option>Pidgin</option><option>Twi</option><option>Hausa</option></select>
   <input type="file" id="fileInput" hidden onchange="handleFile()" accept=".pdf,.docx,.txt,.csv,.json,.py,.js,.html,.md,.jpg,.jpeg,.png,.webp">
   <button class="btn btn2" onclick="document.getElementById('fileInput').click()">📎</button>
   <input id="q" placeholder="Ask, upload a file, or speak" autocomplete="off">
   <button class="btn" onclick="send()">➤</button>
   <button class="btn btn2" id="micBtn" onclick="toggleMic()">🎤</button>
   <button class="btn btn2" id="spkBtn" onclick="toggleSpeak()" title="Read answers aloud">🔊</button>
  </div>
  <audio id="player" controls style="display:none;width:100%;max-width:900px;margin:8px auto 0"></audio>
 </div>
</div>
<script>
const LOGGED_IN = {{ 'true' if user_email else 'false' }};
let currentConv=null, guestHist=[], attachText='', attachImg='', attachName='', busy=false;
let speak = localStorage.getItem('morr_speak')!=='0';
let recorder, chunks=[], recording=false, recMime='';
const $=id=>document.getElementById(id);
function esc(s){return String(s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;');}
function fmt(s){return esc(s).replace(/\*\*(.+?)\*\*/g,'<b>$1</b>').replace(/`([^`\n]+)`/g,'<code>$1</code>').replace(/(https?:\/\/[^\s<]+)/g,'<a href="$1" target="_blank" rel="noopener">$1</a>');}
function addMsg(role,html){const d=document.createElement('div');d.className='msg '+role;d.innerHTML=html;$('chat').appendChild(d);$('chat').scrollTop=$('chat').scrollHeight;return d;}
function srcHtml(srcs){if(!srcs||!srcs.length)return '';return '<span class="src">Sources: '+srcs.slice(0,3).map(s=>'<a href="'+esc(s.url)+'" target="_blank" rel="noopener">'+esc((s.title||s.url).substring(0,40))+'</a>').join(' · ')+'</span>';}
function welcome(){addMsg('ai','Akwaaba! I am MORR AI 🇬🇭<br>I speak English, Pidgin, Twi, Hausa and more. I can read your files and photos, find real photos (try: photo of Independence Arch Accra), draw images (try: generate an image of a sunset in Accra) and make files (try: create a PDF about healthy eating). How may I help you today?');}
function toggleSide(){$('sidebar').classList.toggle('open');$('overlay').classList.toggle('open');}
function toggleSpeak(){speak=!speak;localStorage.setItem('morr_speak',speak?'1':'0');$('spkBtn').style.opacity=speak?1:.4;}
$('spkBtn').style.opacity=speak?1:.4;
function grp(iso){const d=new Date(iso),n=new Date();const sd=x=>new Date(x.getFullYear(),x.getMonth(),x.getDate());const diff=Math.round((sd(n)-sd(d))/86400000);if(diff<=0)return'Today';if(diff==1)return'Yesterday';if(diff<=7)return'Previous 7 days';return'Older';}
async function loadConvs(){
 const box=$('histList');box.innerHTML='';
 if(!LOGGED_IN)return;
 try{
  const list=await (await fetch('/api/conversations')).json();
  let last='';
  list.forEach(c=>{
   const g=grp(c.updated);
   if(g!==last){const h=document.createElement('div');h.className='grp';h.textContent=g;box.appendChild(h);last=g;}
   const row=document.createElement('div');row.className='hist'+(c.id===currentConv?' active':'');
   const t=document.createElement('span');t.textContent=c.title||'New chat';
   const x=document.createElement('i');x.textContent='✕';x.title='Delete';
   row.onclick=()=>openConv(c.id);
   x.onclick=ev=>{ev.stopPropagation();delConv(c.id);};
   row.appendChild(t);row.appendChild(x);box.appendChild(row);
  });
 }catch(e){}
}
async function openConv(id){
 try{
  const msgs=await (await fetch('/api/conversations/'+id)).json();
  $('chat').innerHTML='';currentConv=id;
  msgs.forEach(m=>addMsg(m.role==='user'?'user':'ai',fmt(m.content)));
  loadConvs();
  if(window.innerWidth<=800&&$('sidebar').classList.contains('open'))toggleSide();
 }catch(e){$('status').innerText='Could not open that chat';}
}
async function delConv(id){
 if(!confirm('Delete this chat?'))return;
 await fetch('/api/conversations/'+id,{method:'DELETE'});
 if(currentConv===id)newChat();else loadConvs();
}
function newChat(){currentConv=null;guestHist=[];$('chat').innerHTML='';welcome();loadConvs();if(window.innerWidth<=800&&$('sidebar').classList.contains('open'))toggleSide();}

/* Sends a form (file or audio) with upload progress and a friendly error if the server is slow or restarts */
function postForm(url,fd,onProg,timeoutMs){
 return new Promise((resolve,reject)=>{
  const x=new XMLHttpRequest();x.open('POST',url);x.timeout=timeoutMs||120000;
  x.upload.onprogress=e=>{if(e.lengthComputable&&onProg)onProg(Math.round(e.loaded/e.total*100));};
  x.onload=()=>{try{resolve(JSON.parse(x.responseText));}catch(e){reject(new Error('The server is busy or restarting ('+x.status+'). Please try again in a moment.'));}};
  x.onerror=()=>reject(new Error('Network problem. Check your internet and try again.'));
  x.ontimeout=()=>reject(new Error('This is taking too long. Please try again.'));
  x.send(fd);
 });
}
/* Shrinks big photos in the browser before upload (a 4 MB phone photo becomes about 200 KB) */
function compressImage(file,maxDim,quality){
 return new Promise(resolve=>{
  if(!/^image\/(jpeg|png|webp)$/.test(file.type)||file.size<250*1024){resolve(file);return;}
  const img=new Image();const url=URL.createObjectURL(file);
  img.onload=()=>{
   const s=Math.min(1,maxDim/Math.max(img.width,img.height));
   const w=Math.max(1,Math.round(img.width*s)),h=Math.max(1,Math.round(img.height*s));
   const c=document.createElement('canvas');c.width=w;c.height=h;const ctx=c.getContext('2d');
   ctx.fillStyle='#fff';ctx.fillRect(0,0,w,h);ctx.drawImage(img,0,0,w,h);
   c.toBlob(b=>{URL.revokeObjectURL(url);
    if(b&&b.size<file.size)resolve(new File([b],file.name.replace(/\.\w+$/,'')+'.jpg',{type:'image/jpeg'}));else resolve(file);
   },'image/jpeg',quality);
  };
  img.onerror=()=>{URL.revokeObjectURL(url);resolve(file);};
  img.src=url;
 });
}
async function send(){
 if(busy)return;
 const i=$('q');const q=i.value.trim();
 if(!q&&!attachText&&!attachImg)return;
 busy=true;
 if(q)addMsg('user',fmt(q));
 i.value='';$('status').innerText='MORR is thinking...';
 const body={question:q,file_context:attachText,image_context:attachImg,attachment_name:attachName,language:$('lang').value,
  conversation_id:currentConv,history:LOGGED_IN?[]:guestHist,speak:speak};
 attachText='';attachImg='';attachName='';$('fileInput').value='';
 try{
  const ctl=new AbortController();const tm=setTimeout(()=>ctl.abort(),150000);
  const res=await fetch('/ask',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(body),signal:ctl.signal});
  clearTimeout(tm);
  let data;try{data=await res.json();}catch(e){throw new Error('The server is busy or restarting. Please try again in a moment.');}
  $('status').innerText='';
  if(data.limit){addMsg('ai',fmt(data.answer)+'<br><a href="'+esc(data.upgrade_url)+'">'+(data.upgrade_url==='/signup'?'Create a free account':'⭐ See Premium')+'</a>');busy=false;loadUsage();return;}
  let extra='';
  if(data.image_url)extra+='<br><img src="'+esc(data.image_url)+'"><br><a href="'+esc(data.download_url||data.image_url)+'">⬇️ Download image</a>';
  if(data.file_url)extra+='<br><a href="'+esc(data.file_url)+'">⬇️ Download '+esc(data.file_name||'file')+'</a>';
  addMsg('ai',fmt(data.answer||'No answer')+extra+srcHtml(data.sources));
  if(LOGGED_IN&&data.conversation_id){currentConv=data.conversation_id;loadConvs();}
  if(!LOGGED_IN){guestHist.push({role:'user',content:q},{role:'assistant',content:data.answer||''});guestHist=guestHist.slice(-10);}
  if(data.audio_url){const p=$('player');p.style.display='block';p.src=data.audio_url;p.play().catch(()=>{});}
 }catch(e){$('status').innerText='';addMsg('ai','Error: '+esc(e.name==='AbortError'?'That took too long. Please try again.':e.message));}
 busy=false;loadUsage();
}
async function handleFile(){
 let f=$('fileInput').files[0];if(!f)return;
 const origName=f.name;
 try{
  let localPrev=null;
  if(/^image\//.test(f.type)){
   $('status').innerText='Preparing '+origName+'...';
   localPrev=URL.createObjectURL(f);
   f=await compressImage(f,1280,0.8);
  }
  const kb=Math.max(1,Math.round(f.size/1024));
  $('status').innerText='Uploading '+origName+' ('+kb+' KB)...';
  const fd=new FormData();fd.append('file',f);
  const data=await postForm('/upload',fd,p=>{$('status').innerText=p<100?('Uploading '+origName+'... '+p+'%'):('Reading '+origName+'... this can take a few seconds');},150000);
  if(data.error){$('status').innerText='';addMsg('ai','⚠️ '+esc(data.error)+(data.upgrade_url?'<br><a href="'+esc(data.upgrade_url)+'">'+(data.upgrade_url==='/signup'?'Create a free account':'⭐ See Premium')+'</a>':''));loadUsage();attachText='';attachImg='';attachName='';return;}
  attachText=data.text||'';attachImg=data.image_desc||'';attachName=origName;
  $('status').innerText='✅ '+origName+' attached. Now type your question (or just press send).';loadUsage();
  const pv=data.preview||localPrev;
  if(pv)addMsg('user','📎 '+esc(origName)+'<br><img src="'+pv+'">');
  else addMsg('user','📎 '+esc(origName));
 }catch(e){attachText='';attachImg='';attachName='';$('status').innerText='⚠️ Upload failed: '+e.message;}
}
async function toggleMic(){
 const btn=$('micBtn');
 try{
  if(!recording){
   if(!navigator.mediaDevices||!window.MediaRecorder){$('status').innerText='Voice recording is not supported in this browser. Try Chrome or Safari.';return;}
   const s=await navigator.mediaDevices.getUserMedia({audio:true});
   const pref=['audio/webm;codecs=opus','audio/webm','audio/mp4','audio/ogg'];
   recMime=pref.find(t=>MediaRecorder.isTypeSupported(t))||'';
   recorder=recMime?new MediaRecorder(s,{mimeType:recMime}):new MediaRecorder(s);chunks=[];
   recorder.ondataavailable=e=>{if(e.data&&e.data.size)chunks.push(e.data);};
   recorder.onstop=async()=>{
    s.getTracks().forEach(t=>t.stop());
    const type=(recorder.mimeType||recMime||'audio/webm').split(';')[0];
    const ext=type.includes('mp4')?'m4a':(type.includes('ogg')?'ogg':'webm');
    const blob=new Blob(chunks,{type});
    if(blob.size<1500){$('status').innerText='That was too short. Tap the mic, speak, then tap again to stop.';return;}
    const fd=new FormData();fd.append('audio',blob,'voice.'+ext);fd.append('language',$('lang').value);
    $('status').innerText='Transcribing...';
    try{
     const d=await postForm('/voice',fd,null,60000);
     if(d.text){$('q').value=d.text;$('status').innerText='';send();}
     else{$('status').innerText=d.error||'Could not understand the audio, please try again.';}
    }catch(e){$('status').innerText='Voice failed: '+e.message;}
   };
   recorder.start();recording=true;btn.innerText='🔴';$('status').innerText='Listening... tap 🔴 when you finish';
  }else{recorder.stop();recording=false;btn.innerText='🎤';}
 }catch(e){recording=false;btn.innerText='🎤';$('status').innerText='Mic error: '+e.message;}
}
$('q').addEventListener('keydown',e=>{if(e.key==='Enter')send();});
async function loadUsage(){
 try{
  const u=await (await fetch('/api/usage')).json();const L=u.limits,U=u.used;let h='';
  if(u.plan==='premium'){
   h='<b>⭐ Premium</b><br>'+(u.days_left!=null?u.days_left+' day(s) left'+(u.auto_renew?' · auto-renews':''):'Unlimited access');
   h+=(u.days_left!=null&&u.days_left<=3&&!u.auto_renew)?'<br><a href="/pricing">Renew now</a>':'<br><a href="/pricing">Manage plan</a>';
  }else{
   h='<b>'+(u.plan==='guest'?'Guest':'Free plan')+'</b> · today<br>Messages '+(U.msg||0)+'/'+L.msg;
   if(L.img)h+='<br>Images '+(U.img||0)+'/'+L.img;else h+='<br>Images, files and uploads need a free account';
   if(L.upload)h+='<br>Uploads '+(U.upload||0)+'/'+L.upload;
   if(L.file)h+='<br>Files '+(U.file||0)+'/'+L.file;
   h+=u.plan==='guest'?'<br><a href="/signup">Create free account</a>':'<br><a href="/pricing">⭐ Go Premium</a>';
  }
  $('usageBox').innerHTML=h;
 }catch(e){}
}
welcome();loadConvs();loadUsage();
</script></body></html>"""

AUTH_HTML = r"""<!DOCTYPE html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>MORR AI GH - {{ title }}</title></head>
<body style="margin:0;background:#0a0a0a;color:white;font-family:system-ui,sans-serif;min-height:100vh;display:flex;align-items:center;justify-content:center">
<div style="background:#171717;padding:28px;border-radius:15px;width:min(340px,90vw);border:1px solid #333">
<h2 style="color:#FFD700;margin-top:0">MORR AI GH 🇬🇭</h2><div style="margin-bottom:10px">{{ title }}</div>
{% if error %}<div style="background:#3a1515;border:1px solid #7a2a2a;padding:8px;border-radius:8px;font-size:13px;margin-bottom:8px">{{ error }}</div>{% endif %}
<form method="POST">
<input name="email" type="email" placeholder="Email" value="{{ email or '' }}" required style="width:100%;padding:12px;margin:6px 0;border-radius:8px;border:1px solid #333;background:#222;color:white;font-size:16px">
<input name="password" type="password" placeholder="Password{% if mode == 'signup' %} (min 8 characters){% endif %}" required style="width:100%;padding:12px;margin:6px 0;border-radius:8px;border:1px solid #333;background:#222;color:white;font-size:16px">
<button style="width:100%;padding:12px;border-radius:8px;border:none;background:#FFD700;color:black;font-weight:bold;margin-top:8px;font-size:16px">{{ title }}</button>
</form>
<div style="margin-top:14px;font-size:13px;color:#999">
{% if mode == 'signup' %}Already have an account? <a href="/login" style="color:#FFD700">Log in</a>{% else %}New here? <a href="/signup" style="color:#FFD700">Create an account</a>{% endif %}
 · <a href="/" style="color:#999">Back</a></div></div></body></html>"""

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def current_user():
    return session.get("user_id")


@app.route("/")
def home():
    return render_template_string(HTML, user_email=session.get("email"))


@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "POST":
        email = (request.form.get("email") or "").lower().strip()
        pw = request.form.get("password") or ""
        err = None
        ipkey = "ip:" + hash_key(client_ip())
        try:
            too_many = get_usage(ipkey).get("signup", 0) >= SIGNUPS_PER_IP
        except Exception:
            too_many = False
        if too_many:
            err = "Too many new accounts were created from this network today. Please try again tomorrow."
        elif not EMAIL_RE.match(email):
            err = "Please enter a valid email."
        elif len(pw) < 8:
            err = "Password must be at least 8 characters."
        if not err:
            try:
                if db("SELECT id FROM users WHERE email=?", (email,), fetch="one"):
                    err = "An account with that email already exists. Please log in."
                else:
                    uid = db("INSERT INTO users (email, password) VALUES (?, ?)",
                             (email, generate_password_hash(pw)), insert=True)
                    add_usage(ipkey, "signup")
                    session.clear()
                    session.permanent = True
                    session["user_id"], session["email"] = uid, email
                    return redirect("/")
            except Exception as e:
                print(f"[signup] failed: {e}")
                err = "Could not create the account right now. Please try again."
        return render_template_string(AUTH_HTML, title="Create account", mode="signup", error=err, email=email)
    return render_template_string(AUTH_HTML, title="Create account", mode="signup", error=None, email="")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = (request.form.get("email") or "").lower().strip()
        pw = request.form.get("password") or ""
        lockkey = "ip:" + hash_key(client_ip())
        try:
            locked = get_usage(lockkey).get("login_fail", 0) >= 15
        except Exception:
            locked = False
        if locked:
            return render_template_string(AUTH_HTML, title="Log in", mode="login",
                                          error="Too many failed attempts. Please try again tomorrow.", email=email)
        try:
            row = db("SELECT id, password FROM users WHERE email=?", (email,), fetch="one")
        except Exception as e:
            print(f"[login] failed: {e}")
            return render_template_string(AUTH_HTML, title="Log in", mode="login",
                                          error="Service busy, please try again.", email=email)
        if row and check_password_hash(row[1], pw):
            session.clear()
            session.permanent = True
            session["user_id"], session["email"] = row[0], email
            return redirect("/")
        add_usage(lockkey, "login_fail")
        return render_template_string(AUTH_HTML, title="Log in", mode="login",
                                      error="Incorrect email or password.", email=email)
    return render_template_string(AUTH_HTML, title="Log in", mode="login", error=None, email="")


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/")


# ===========================================================================
# CHAT HISTORY API (private to each user)
# ===========================================================================
def owns(cid, uid):
    try:
        cid = int(cid)
    except (TypeError, ValueError):
        return None
    return db("SELECT id FROM conversations WHERE id=? AND user_id=?", (cid, uid), fetch="one")


@app.route("/api/conversations")
def api_conversations():
    uid = current_user()
    if not uid:
        return jsonify([]), 401
    rows = db("SELECT id, title, updated FROM conversations WHERE user_id=? ORDER BY updated DESC LIMIT 100",
              (uid,), fetch="all") or []
    return jsonify([{"id": r[0], "title": r[1], "updated": r[2]} for r in rows])


@app.route("/api/conversations/<int:cid>", methods=["GET", "DELETE"])
def api_conversation(cid):
    uid = current_user()
    if not uid:
        return jsonify({"error": "login required"}), 401
    if not owns(cid, uid):
        return jsonify({"error": "not found"}), 404
    if request.method == "DELETE":
        db("DELETE FROM messages WHERE conversation_id=?", (cid,))
        db("DELETE FROM conversations WHERE id=?", (cid,))
        return jsonify({"ok": True})
    rows = db("SELECT role, content FROM messages WHERE conversation_id=? ORDER BY id", (cid,), fetch="all") or []
    return jsonify([{"role": r[0], "content": r[1]} for r in rows])


# ===========================================================================
# PLANS, DAILY LIMITS AND PAYMENTS (Paystack)
# ===========================================================================
def today():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")  # Ghana is on UTC, so this resets at midnight Ghana time


def parse_iso(s):
    try:
        return datetime.fromisoformat(s) if s else None
    except Exception:
        return None


def client_ip():
    return request.remote_addr or "unknown"


def hash_key(s):
    return hashlib.sha256((s + SECRET).encode()).hexdigest()[:20]


def usage_key(uid):
    return f"u:{uid}" if uid else "g:" + hash_key(client_ip())


def is_admin():
    return (session.get("email") or "").lower() in ADMIN_EMAILS


def get_premium_until(uid):
    row = db("SELECT premium_until FROM subscriptions WHERE user_id=?", (uid,), fetch="one")
    return row[0] if row and row[0] else None


def get_plan(uid, email=None):
    if not uid:
        return "guest"
    if email and email.lower() in ADMIN_EMAILS:
        return "premium"
    try:
        until = parse_iso(get_premium_until(uid))
        if until and until > datetime.now(timezone.utc):
            return "premium"
    except Exception as e:
        print(f"[plan] lookup failed: {str(e)[:150]}")
    return "free"


def premium_info(uid):
    info = {"days_left": None, "until": None, "auto_renew": False}
    try:
        row = db("SELECT premium_until, auto_renew FROM subscriptions WHERE user_id=?", (uid,), fetch="one")
        if row:
            until = parse_iso(row[0])
            now = datetime.now(timezone.utc)
            if until and until > now:
                info["days_left"] = max(1, int((until - now).total_seconds() // 86400) + 1)
                info["until"] = until.strftime("%d %b %Y")
            info["auto_renew"] = bool(row[1])
    except Exception as e:
        print(f"[plan] info failed: {str(e)[:150]}")
    return info


def get_usage(key):
    rows = db("SELECT kind, n FROM usage_daily WHERE user_key=? AND day=?", (key, today()), fetch="all") or []
    return {k: n for k, n in rows}


def add_usage(key, kind, amount=1):
    try:
        db("INSERT INTO usage_daily (user_key, day, kind, n) VALUES (?, ?, ?, ?) "
           "ON CONFLICT (user_key, day, kind) DO UPDATE SET n = usage_daily.n + ?",
           (key, today(), kind, amount, amount))
    except Exception as e:
        print(f"[usage] could not record: {str(e)[:150]}")


def check_limit(uid, email, kind):
    """Returns (allowed, plan, used, limit). If the database hiccups we let the user through."""
    plan = get_plan(uid, email)
    limit = LIMITS[plan][kind]
    try:
        used = get_usage(usage_key(uid)).get(kind, 0)
    except Exception as e:
        print(f"[usage] lookup failed: {str(e)[:150]}")
        return True, plan, 0, limit
    return used < limit, plan, used, limit


KIND_NAME = {"msg": "messages", "img": "image generations", "upload": "file uploads", "voice": "voice messages",
             "file": "file creations"}


def limit_response(plan, kind, limit):
    name = KIND_NAME[kind]
    if plan == "guest":
        if limit == 0:
            text = f"{name.capitalize()} need a free account. Please sign up or log in to use them."
        else:
            text = f"You have used all {limit} guest {name} for today. Create a free account to get more."
        url = "/signup"
    elif plan == "free":
        text = (f"You have used all {limit} free {name} for today. Upgrade to Premium for unlimited use, "
                "or come back tomorrow (it resets at midnight).")
        url = "/pricing"
    else:
        text = f"You have reached today's fair-use limit for {name}. It resets at midnight, please try again tomorrow."
        url = "/pricing"
    return {"answer": text, "error": text, "limit": True, "upgrade_url": url}


def set_premium(uid, days, mode="add"):
    """add: extend from the later of now / current end. floor: make sure it lasts at least `days` from now.
    end: finish premium immediately."""
    now = datetime.now(timezone.utc)
    cur = parse_iso(get_premium_until(uid))
    if mode == "add":
        until = (cur if (cur and cur > now) else now) + timedelta(days=days)
    elif mode == "floor":
        until = max(cur, now + timedelta(days=days)) if cur else now + timedelta(days=days)
    else:
        until = now
    ts = until.isoformat()
    db("INSERT INTO subscriptions (user_id, premium_until, updated) VALUES (?, ?, ?) "
       "ON CONFLICT (user_id) DO UPDATE SET premium_until = ?, updated = ?",
       (uid, ts, now_iso(), ts, now_iso()))
    return ts


def save_sub_fields(uid, code, token, auto):
    db("INSERT INTO subscriptions (user_id, sub_code, email_token, auto_renew, updated) VALUES (?, ?, ?, ?, ?) "
       "ON CONFLICT (user_id) DO UPDATE SET sub_code = ?, email_token = ?, auto_renew = ?, updated = ?",
       (uid, code, token, 1 if auto else 0, now_iso(), code, token, 1 if auto else 0, now_iso()))


def user_id_by_email(email):
    row = db("SELECT id FROM users WHERE email=?", ((email or "").lower().strip(),), fetch="one")
    return row[0] if row else None


def paystack(method, path, **kw):
    if not PAYSTACK_SECRET:
        raise RuntimeError("Paystack is not configured")
    r = requests.request(method, "https://api.paystack.co" + path,
                         headers={"Authorization": f"Bearer {PAYSTACK_SECRET}"}, timeout=20, **kw)
    try:
        j = r.json()
    except Exception:
        j = {}
    if r.status_code >= 400 or not j.get("status"):
        raise RuntimeError(j.get("message") or f"HTTP {r.status_code}")
    return j["data"]


def apply_paystack_payment(data):
    """data is a Paystack transaction object that has been verified (by API call or signed webhook).
    Safe to call more than once for the same payment. Returns (ok, message)."""
    if data.get("status") != "success":
        return False, "payment not successful"
    if (data.get("currency") or "GHS").upper() != "GHS" or int(data.get("amount") or 0) < int(round(PRICE_GHS * 100)):
        print(f"[pay] ignoring payment with wrong amount/currency: {data.get('amount')} {data.get('currency')}")
        return False, "wrong amount"
    ref = data.get("reference")
    if not ref:
        return False, "no reference"
    meta = data.get("metadata")
    if not isinstance(meta, dict):
        meta = {}
    email = ((data.get("customer") or {}).get("email") or "").lower()
    uid = None
    try:
        uid = int(meta.get("user_id")) if meta.get("user_id") else None
    except (TypeError, ValueError):
        uid = None
    if uid and not db("SELECT id FROM users WHERE id=?", (uid,), fetch="one"):
        uid = None
    if not uid:
        uid = user_id_by_email(email)
    if not uid:
        print(f"[pay] no user found for payment {ref}")
        return False, "user not found"
    if db("SELECT id FROM payments WHERE reference=?", (ref,), fetch="one"):
        return True, "already processed"
    is_sub = bool(data.get("plan")) or meta.get("kind") == "subscription"
    try:
        db("INSERT INTO payments (reference, user_id, email, amount, kind, status, created) VALUES (?, ?, ?, ?, ?, ?, ?)",
           (ref, uid, email, int(data.get("amount") or 0), "subscription" if is_sub else "onetime", "success", now_iso()))
    except Exception:
        if db("SELECT id FROM payments WHERE reference=?", (ref,), fetch="one"):
            return True, "already processed"
        raise
    if is_sub:
        set_premium(uid, DAYS_SUBSCRIPTION, "floor")
    else:
        set_premium(uid, DAYS_ONE_TIME, "add")
    print(f"[pay] premium activated for user {uid} ({'subscription' if is_sub else 'one-time'})")
    return True, "activated"


@app.route("/api/usage")
def api_usage():
    uid, email = current_user(), session.get("email")
    plan = get_plan(uid, email)
    try:
        used = get_usage(usage_key(uid))
    except Exception:
        used = {}
    info = premium_info(uid) if uid else {"days_left": None, "until": None, "auto_renew": False}
    if plan == "premium" and (email or "").lower() in ADMIN_EMAILS and not info["days_left"]:
        info["days_left"] = None
    return jsonify({"plan": plan, "used": used, "limits": LIMITS[plan], **info})


@app.route("/pricing")
def pricing():
    uid, email = current_user(), session.get("email")
    plan = get_plan(uid, email)
    info = premium_info(uid) if uid else {"days_left": None, "until": None, "auto_renew": False}
    return render_template_string(
        PRICING_HTML, plan=plan, logged_in=bool(uid), email=email, price=PRICE_GHS, limits=LIMITS,
        free_mb=MAX_UPLOAD_MB["free"], prem_mb=MAX_UPLOAD_MB["premium"], info=info,
        pay_ready=bool(PAYSTACK_SECRET), plan_ready=bool(PAYSTACK_PLAN),
        status=request.args.get("status"), error=request.args.get("error"))


@app.route("/billing/checkout", methods=["POST"])
def billing_checkout():
    uid = current_user()
    if not uid:
        return redirect("/login")
    kind = request.form.get("kind")
    if kind not in ("subscription", "onetime"):
        return redirect("/pricing?error=Please choose a payment option.")
    if not PAYSTACK_SECRET or (kind == "subscription" and not PAYSTACK_PLAN):
        return redirect("/pricing?error=Payments are not switched on yet. Please try again later.")
    base = SITE_URL or request.host_url.rstrip("/")
    payload = {"email": session["email"], "amount": int(round(PRICE_GHS * 100)), "currency": "GHS",
               "callback_url": base + "/billing/callback", "metadata": {"user_id": uid, "kind": kind}}
    if kind == "subscription":
        payload["plan"] = PAYSTACK_PLAN
        payload["channels"] = ["card"]
    else:
        payload["channels"] = ["mobile_money", "card"]
    try:
        data = paystack("POST", "/transaction/initialize", json=payload)
        return redirect(data["authorization_url"])
    except Exception as e:
        print(f"[pay] initialize failed: {str(e)[:200]}")
        return redirect("/pricing?error=We could not start the payment. Please try again.")


@app.route("/billing/callback")
def billing_callback():
    ref = request.args.get("reference") or request.args.get("trxref")
    if not ref:
        return redirect("/pricing")
    try:
        data = paystack("GET", "/transaction/verify/" + urllib.parse.quote(ref, safe=""))
        ok, msg = apply_paystack_payment(data)
        return redirect("/pricing?status=success" if ok else "/pricing?status=failed")
    except Exception as e:
        print(f"[pay] verify failed: {str(e)[:200]}")
        return redirect("/pricing?status=failed")


@app.route("/paystack/webhook", methods=["POST"])
def paystack_webhook():
    raw = request.get_data()
    sig = request.headers.get("x-paystack-signature", "")
    if not PAYSTACK_SECRET:
        return "not configured", 503
    expected = hmac.new(PAYSTACK_SECRET.encode(), raw, hashlib.sha512).hexdigest()
    if not hmac.compare_digest(sig, expected):
        return "bad signature", 401
    try:
        event = json.loads(raw)
        typ, data = event.get("event"), event.get("data") or {}
        if typ == "charge.success":
            apply_paystack_payment(data)
        elif typ in ("subscription.create", "subscription.disable", "subscription.not_renew"):
            uid = user_id_by_email((data.get("customer") or {}).get("email"))
            if uid:
                if typ == "subscription.create":
                    save_sub_fields(uid, data.get("subscription_code"), data.get("email_token"), True)
                else:
                    row = db("SELECT sub_code, email_token FROM subscriptions WHERE user_id=?", (uid,), fetch="one")
                    save_sub_fields(uid, row[0] if row else None, row[1] if row else None, False)
    except Exception as e:
        print(f"[pay] webhook failed: {str(e)[:200]}")
        return "error", 500  # Paystack will retry
    return "ok", 200


@app.route("/billing/cancel", methods=["POST"])
def billing_cancel():
    uid = current_user()
    if not uid:
        return redirect("/login")
    try:
        row = db("SELECT sub_code, email_token FROM subscriptions WHERE user_id=?", (uid,), fetch="one")
        if not row or not row[0] or not row[1]:
            return redirect("/pricing?error=No active card subscription was found on your account.")
        paystack("POST", "/subscription/disable", json={"code": row[0], "token": row[1]})
        save_sub_fields(uid, row[0], row[1], False)
        return redirect("/pricing?status=cancelled")
    except Exception as e:
        print(f"[pay] cancel failed: {str(e)[:200]}")
        return redirect("/pricing?error=We could not cancel right now. Please try again.")


# ---------------------------------------------------------------------------
# ADMIN (only emails listed in ADMIN_EMAILS)
# ---------------------------------------------------------------------------
@app.route("/admin")
def admin_page():
    if not is_admin():
        return "Not allowed", 403
    users_total = db("SELECT COUNT(*) FROM users", fetch="one")[0]
    now = datetime.now(timezone.utc)
    prem = 0
    for (u,) in db("SELECT premium_until FROM subscriptions", fetch="all") or []:
        d = parse_iso(u)
        if d and d > now:
            prem += 1
    msgs = db("SELECT COALESCE(SUM(n), 0) FROM usage_daily WHERE day=? AND kind='msg'", (today(),), fetch="one")[0]
    revenue = db("SELECT COALESCE(SUM(amount), 0) FROM payments", fetch="one")[0]
    pays = db("SELECT email, amount, kind, created FROM payments ORDER BY id DESC LIMIT 15", fetch="all") or []
    return render_template_string(ADMIN_HTML, users_total=users_total, prem=prem, msgs=msgs,
                                  revenue=revenue / 100.0, pays=pays, msg=request.args.get("msg"))


@app.route("/admin/grant", methods=["POST"])
def admin_grant():
    if not is_admin():
        return "Not allowed", 403
    email = (request.form.get("email") or "").lower().strip()
    try:
        days = max(1, min(int(request.form.get("days") or 30), 3650))
    except ValueError:
        days = 30
    uid = user_id_by_email(email)
    if not uid:
        return redirect("/admin?msg=No user with that email.")
    set_premium(uid, days, "add")
    return redirect(f"/admin?msg=Gave {days} days of Premium to {email}.")


@app.route("/admin/revoke", methods=["POST"])
def admin_revoke():
    if not is_admin():
        return "Not allowed", 403
    email = (request.form.get("email") or "").lower().strip()
    uid = user_id_by_email(email)
    if not uid:
        return redirect("/admin?msg=No user with that email.")
    set_premium(uid, 0, "end")
    return redirect(f"/admin?msg=Premium ended for {email}.")


# ---------------------------------------------------------------------------
# PRICING / ADMIN / LEGAL PAGES
# ---------------------------------------------------------------------------
PAGE_STYLE = ("body{margin:0;background:#0a0a0a;color:white;font-family:system-ui,sans-serif;padding:20px}"
              ".wrap{max-width:860px;margin:0 auto}a{color:#FFD700}h1{color:#FFD700;margin:6px 0}"
              ".cards{display:flex;gap:14px;flex-wrap:wrap;margin-top:16px}"
              ".card{flex:1;min-width:260px;background:#171717;border:1px solid #333;border-radius:14px;padding:18px}"
              ".card.prem{border-color:#FFD700}.price{font-size:28px;color:#FFD700;font-weight:bold}"
              "ul{padding-left:18px;line-height:1.8;color:#ddd}"
              "button{width:100%;padding:12px;border-radius:10px;border:none;background:#FFD700;color:black;font-weight:bold;"
              "font-size:15px;cursor:pointer;margin-top:8px}button.alt{background:#333;color:white}button:disabled{opacity:.4}"
              ".ok{background:#12301a;border:1px solid #2a7a3a;padding:10px;border-radius:10px;margin:10px 0}"
              ".err{background:#3a1515;border:1px solid #7a2a2a;padding:10px;border-radius:10px;margin:10px 0}"
              "table{width:100%;border-collapse:collapse}td,th{padding:8px;border-bottom:1px solid #2a2a2a;text-align:left;font-size:14px}"
              "input{padding:10px;border-radius:8px;border:1px solid #333;background:#222;color:white;font-size:15px}")

PRICING_HTML = r"""<!DOCTYPE html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>MORR AI GH - Plans</title>
<style>""" + PAGE_STYLE + r"""</style></head><body><div class="wrap">
<a href="/">← Back to chat</a>
<h1>MORR AI GH Plans</h1>
{% if status == 'success' %}<div class="ok">🎉 Payment received. Premium is now active. Thank you!</div>
{% elif status == 'failed' %}<div class="err">We could not confirm your payment yet. If you were charged, Premium will switch on automatically in a few minutes. If not, please contact support.</div>
{% elif status == 'cancelled' %}<div class="ok">Auto-renewal cancelled. You keep Premium until the end of the period you paid for.</div>{% endif %}
{% if error %}<div class="err">{{ error }}</div>{% endif %}
{% if plan == 'premium' %}<div class="ok">⭐ You are on Premium{% if info.until %} until {{ info.until }}{% if info.auto_renew %} (auto-renews){% endif %}{% endif %}.</div>{% endif %}
<div class="cards">
 <div class="card"><h2>Free</h2><div class="price">GHS 0</div>
  <ul><li>{{ limits.free.msg }} messages a day</li><li>{{ limits.free.img }} AI image generations a day</li>
  <li>{{ limits.free.file }} PDF, Word or Excel file creations a day</li>
  <li>{{ limits.free.upload }} file or photo uploads a day (up to {{ free_mb }} MB)</li><li>Saved chat history</li></ul>
  {% if not logged_in %}<a href="/signup"><button class="alt" type="button">Create free account</button></a>{% endif %}</div>
 <div class="card prem"><h2>⭐ Premium</h2><div class="price">GHS {{ '%.2f'|format(price) }} <span style="font-size:14px;color:#aaa">/ month</span></div>
  <ul><li>Unlimited messages*</li><li>Unlimited AI image generations*</li><li>Unlimited PDF, Word and Excel creation*</li><li>Unlimited file and photo uploads* (up to {{ prem_mb }} MB)</li>
  <li>Priority access to our best AI model</li><li>Saved chat history</li></ul>
  <div style="font-size:12px;color:#888">*Fair use applies.</div>
  {% if not logged_in %}<a href="/login"><button type="button">Log in to upgrade</button></a>
  {% else %}
  <form method="POST" action="/billing/checkout"><input type="hidden" name="kind" value="subscription">
   <button {% if not (pay_ready and plan_ready) %}disabled{% endif %}>Subscribe monthly with card (auto-renews)</button></form>
  <form method="POST" action="/billing/checkout"><input type="hidden" name="kind" value="onetime">
   <button class="alt" {% if not pay_ready %}disabled{% endif %}>Pay for 30 days (Mobile Money or card)</button></form>
  {% if plan == 'premium' and info.auto_renew %}<form method="POST" action="/billing/cancel"><button class="alt" onclick="return confirm('Stop automatic renewal?')">Cancel auto-renewal</button></form>{% endif %}
  {% if not pay_ready %}<div style="font-size:12px;color:#c88;margin-top:8px">Payments are not switched on yet.</div>{% endif %}
  {% endif %}</div>
</div>
<p style="color:#888;font-size:13px;margin-top:20px">Payments are processed securely by Paystack. We never see or store your card or Mobile Money details.
<a href="/terms">Terms</a> · <a href="/privacy">Privacy</a></p>
</div></body></html>"""

ADMIN_HTML = r"""<!DOCTYPE html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>MORR AI GH - Admin</title>
<style>""" + PAGE_STYLE + r"""</style></head><body><div class="wrap">
<a href="/">← Back to chat</a><h1>Admin</h1>
{% if msg %}<div class="ok">{{ msg }}</div>{% endif %}
<div class="cards"><div class="card"><div>Users</div><div class="price">{{ users_total }}</div></div>
<div class="card"><div>Active Premium</div><div class="price">{{ prem }}</div></div>
<div class="card"><div>Messages today</div><div class="price">{{ msgs }}</div></div>
<div class="card"><div>Total received (GHS)</div><div class="price">{{ '%.2f'|format(revenue) }}</div></div></div>
<h3>Give or end Premium</h3>
<form method="POST" action="/admin/grant" style="display:flex;gap:8px;flex-wrap:wrap"><input name="email" type="email" placeholder="User email" required>
<input name="days" type="number" value="30" min="1" style="width:90px"><button style="width:auto;padding:10px 16px">Give Premium</button></form>
<form method="POST" action="/admin/revoke" style="display:flex;gap:8px;margin-top:8px;flex-wrap:wrap"><input name="email" type="email" placeholder="User email" required>
<button class="alt" style="width:auto;padding:10px 16px">End Premium</button></form>
<h3>Recent payments</h3><table><tr><th>Email</th><th>GHS</th><th>Type</th><th>When</th></tr>
{% for e, a, k, c in pays %}<tr><td>{{ e }}</td><td>{{ '%.2f'|format(a / 100) }}</td><td>{{ k }}</td><td>{{ c[:16] }}</td></tr>{% endfor %}</table>
</div></body></html>"""

LEGAL_HTML = r"""<!DOCTYPE html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>MORR AI GH - {{ title }}</title>
<style>""" + PAGE_STYLE + r"""</style></head><body><div class="wrap"><a href="/">← Back to chat</a><h1>{{ title }}</h1>
<div class="err">DRAFT TEMPLATE. Replace everything in [square brackets] and have it reviewed by a qualified lawyer before you rely on it.</div>
{% for h, t in sections %}<h3>{{ h }}</h3><p style="color:#ddd;line-height:1.7">{{ t }}</p>{% endfor %}
</div></body></html>"""

TERMS = [
    ("The service", "MORR AI GH is an AI chat assistant that can answer questions, read files and photos, find photos online, generate images and create documents. AI answers can be wrong, so check important information before relying on it. Photos found online belong to their owners and are shown with their source."),
    ("Accounts", "You must give a real email address and keep your password private. You are responsible for activity on your account."),
    ("Acceptable use", "Do not use MORR AI GH for anything illegal, abusive, harmful or deceptive, or to attack or overload the service. We may limit or close accounts that break these rules."),
    ("Plans and limits", "Free accounts have daily limits that reset at midnight (Ghana time). Premium is GHS 53.00 per month and is described as unlimited, subject to fair use."),
    ("Billing", "Card subscriptions renew automatically every month until you cancel on the Plans page. Mobile Money and one-off payments give 30 days of Premium and do not renew by themselves. Payments are handled by Paystack. [Refund policy: write yours here.]"),
    ("Changes and ending", "We may change features, prices or these terms, and will tell you about important changes. You can stop using the service at any time. [Add how users can delete their account.]"),
    ("Contact", "[Your business name, address and support email.]"),
]
PRIVACY = [
    ("What we collect", "Your email address, a scrambled (hashed) version of your password, your chats, files, images and voice recordings you send for processing, and records of your payments (reference, amount, date). We do not see or store your card or Mobile Money details; Paystack handles those."),
    ("How we use it", "To run your account, answer your questions, save your chat history, enforce plan limits and process payments."),
    ("Who we share it with", "Your messages, and any files, images or voice recordings you send, are processed by our AI and search providers (for example Google Gemini and Groq) and our hosting and database providers. Image searches are sent to Wikipedia and any photo sites we enable. Payments are processed by Paystack. [List any others you add.]"),
    ("How long we keep it", "[Say how long you keep chats and account data, and how users can ask for deletion.]"),
    ("Your rights", "[Explain how users can access, correct or delete their data, in line with Ghana's Data Protection Act. Ask a lawyer to check this and whether you must register with the Data Protection Commission.]"),
    ("Contact", "[Your support email.]"),
]


@app.route("/terms")
def terms_page():
    return render_template_string(LEGAL_HTML, title="Terms of Service", sections=TERMS)


@app.route("/privacy")
def privacy_page():
    return render_template_string(LEGAL_HTML, title="Privacy Policy", sections=PRIVACY)


# ===========================================================================
# ASK
# ===========================================================================
NO_COUNT = {"none", "image-failed", "image-clarify", "file-failed"}


@app.route("/ask", methods=["POST"])
def ask_route():
    d = request.get_json(silent=True) or {}
    q = (d.get("question") or "").strip()[:4000]
    file_ctx = (d.get("file_context") or "")[:10000]
    img_ctx = (d.get("image_context") or "")[:4000]
    att_name = (d.get("attachment_name") or "")[:100]
    lang = d.get("language") or "Auto"
    if lang not in ("Auto", "English", "Pidgin", "Twi", "Hausa"):
        lang = "Auto"
    if not q and (file_ctx or img_ctx):
        q = "Please read the attached file or image and summarize the key points."
    if not q:
        return jsonify({"answer": "Please type a question."}), 400

    uid = current_user()
    email_now = session.get("email")
    has_att = bool(file_ctx or img_ctx)
    is_file = wants_file(q)
    is_img = (not is_file) and (not has_att) and wants_image(q)
    kind = "file" if is_file else "msg"
    allowed, plan, used, limit = check_limit(uid, email_now, kind)
    if not allowed:
        return jsonify(limit_response(plan, kind, limit)), 429
    art_ok = check_limit(uid, email_now, "img")[0] if is_img else True

    conv_id = d.get("conversation_id")
    history = []
    try:
        if uid:
            if conv_id and not owns(conv_id, uid):
                conv_id = None
            if conv_id:
                rows = db("SELECT role, content FROM messages WHERE conversation_id=? ORDER BY id DESC LIMIT 10",
                          (int(conv_id),), fetch="all") or []
                history = clean_history([{"role": r[0], "content": r[1]} for r in reversed(rows)])
        else:
            history = clean_history(d.get("history"))
    except Exception as e:
        print(f"[ask] history load failed: {e}")
        conv_id = None

    # The person talks about an attachment that never arrived: say so clearly instead of letting the AI guess
    if not has_att and not history and ATTACH_REF_RE.search(q):
        return jsonify({"answer": "I do not see an attachment yet. Tap the 📎 button, pick your picture or file, wait "
                                  "for the ✅ message, then ask your question again."})

    order = PROVIDER_ORDER if plan == "premium" else FREE_PROVIDER_ORDER
    image_name = file_name = None
    used_kind = "msg"
    if is_file:
        ans, sources, provider, file_name = make_file_answer(q, lang, file_ctx, img_ctx, history, order)
        used_kind = "file"
    elif is_img:
        info = plan_image(q, history)
        if info["mode"] == "art" and not art_ok:
            return jsonify(limit_response(plan, "img", LIMITS[plan]["img"])), 429
        ans, sources, provider, image_name, used_kind = make_image_answer(q, info, art_ok)
    else:
        try:
            ans, sources, provider = generate_answer(q, lang, file_ctx, img_ctx, history, order)
        except Exception as e:
            print(f"[ask] unexpected failure: {e}")
            return jsonify({"answer": "Something went wrong on my side. Please try again."}), 500
    print(f"[ask] answered by {provider}")
    if provider not in NO_COUNT:
        add_usage(usage_key(uid), used_kind)

    if uid:
        try:
            ts = now_iso()
            if not conv_id:
                conv_id = db("INSERT INTO conversations (user_id, title, created, updated) VALUES (?, ?, ?, ?)",
                             (uid, q[:45], ts, ts), insert=True)
            conv_id = int(conv_id)
            stored_q = (f"📎 {att_name}\n" if att_name else "") + q
            db("INSERT INTO messages (conversation_id, role, content, created) VALUES (?, ?, ?, ?)",
               (conv_id, "user", stored_q, ts))
            db("INSERT INTO messages (conversation_id, role, content, created) VALUES (?, ?, ?, ?)",
               (conv_id, "assistant", ans, now_iso()))
            db("UPDATE conversations SET updated=? WHERE id=?", (now_iso(), conv_id))
        except Exception as e:
            print(f"[ask] saving chat failed: {e}")

    audio_url = make_audio(ans, lang) if d.get("speak") and not image_name and not file_name else None
    return jsonify({"answer": ans, "sources": sources, "audio_url": audio_url, "conversation_id": conv_id,
                    "image_url": f"/media/{image_name}" if image_name else None,
                    "download_url": f"/download/{image_name}" if image_name else None,
                    "file_url": f"/download/{file_name}" if file_name else None,
                    "file_name": friendly_name(file_name) if file_name else None})


# ===========================================================================
# FILE & IMAGE READER
# ===========================================================================
IMG_EXT = {".jpg", ".jpeg", ".png", ".webp"}
DESCRIBE_PROMPT = ("Describe this image in detail. Transcribe ALL visible text exactly (OCR). "
                   "Mention objects, people (without naming strangers), colours, charts or tables and what they show.")


@app.route("/upload", methods=["POST"])
def upload():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "No file received."}), 400
    uid_u = current_user()
    allowed, plan_u, used_u, limit_u = check_limit(uid_u, session.get("email"), "upload")
    if not allowed:
        return jsonify(limit_response(plan_u, "upload", limit_u)), 429
    max_mb = MAX_UPLOAD_MB["premium" if plan_u == "premium" else "free"]
    if request.content_length and request.content_length > max_mb * 1024 * 1024:
        msg = f"That file is too big for your plan (max {max_mb} MB)."
        if plan_u != "premium":
            msg += f" Premium allows files up to {MAX_UPLOAD_MB['premium']} MB."
        return jsonify({"error": msg, "upgrade_url": "/pricing" if plan_u != "premium" else None}), 413
    safe = secure_filename(f.filename) or "upload"
    ext = os.path.splitext(safe.lower())[1]
    tmp = os.path.join(tempfile.gettempdir(), f"{uuid.uuid4().hex[:8]}_{safe}")
    f.save(tmp)
    t0 = time.time()
    text, img_desc, preview = "", "", None
    try:
        if ext in IMG_EXT:
            from PIL import Image, ImageOps
            im = ImageOps.exif_transpose(Image.open(tmp)).convert("RGB")  # fixes sideways phone photos
            big = im.copy()
            big.thumbnail((1600, 1600))
            buf = io.BytesIO()
            big.save(buf, "JPEG", quality=85)
            im.thumbnail((800, 800))
            cleanup_old_files()
            pname = f"prev_{uuid.uuid4().hex[:8]}.jpg"
            im.save(os.path.join(GENERATED_DIR, pname), "JPEG", quality=85)
            preview = f"/media/{pname}"
            if not (gclient or GROQ_KEY):
                return jsonify({"error": "Image reading needs GEMINI_API_KEY or GROQ_API_KEY to be set."})
            try:
                img_desc = clean_answer(describe_image(buf.getvalue()))
            except Exception as e:
                print(f"[upload] image read failed: {str(e)[:300]}")
                return jsonify({"error": "Both of my image readers are busy or at their daily limit. Please try again in a few minutes."})
            if not img_desc:
                return jsonify({"error": "I could not make out anything in that image."})
        elif ext == ".pdf":
            import PyPDF2
            reader = PyPDF2.PdfReader(tmp)
            text = " ".join((p.extract_text() or "") for p in reader.pages[:15]).strip()[:10000]
            if len(text) < 30 and gclient and os.path.getsize(tmp) < 10 * 1024 * 1024:
                with open(tmp, "rb") as fh:
                    data = fh.read()
                try:
                    resp = gemini_generate([types.Content(role="user", parts=[
                        types.Part.from_bytes(data=data, mime_type="application/pdf"),
                        types.Part(text="Extract all the text from this document.")])])
                    text = (resp.text or "")[:10000]
                except Exception as e:
                    print(f"[upload] scanned pdf failed: {str(e)[:300]}")
            if not text:
                return jsonify({"error": "I could not find readable text in that PDF."})
        elif ext in (".txt", ".csv", ".py", ".js", ".html", ".json", ".md"):
            with open(tmp, "r", errors="ignore") as fh:
                text = fh.read()[:10000]
        elif ext == ".docx":
            import docx
            text = " ".join(p.text for p in docx.Document(tmp).paragraphs)[:10000]
        else:
            return jsonify({"error": "That file type is not supported. Try PDF, DOCX, TXT, CSV, JSON, JPG, PNG or WEBP."})
    except Exception as e:
        print(f"[upload] failed: {str(e)[:300]}")
        return jsonify({"error": "I could not read that file. It may be damaged or password protected."})
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    print(f"[upload] {ext} processed in {time.time() - t0:.1f}s")
    add_usage(usage_key(uid_u), "upload")
    return jsonify({"text": text, "image_desc": img_desc, "preview": preview})


@app.route("/voice", methods=["POST"])
def voice():
    f = request.files.get("audio")
    if not f:
        return jsonify({"text": "", "error": "No audio was received. Please try again."})
    uid_v = current_user()
    allowed, plan_v, used_v, limit_v = check_limit(uid_v, session.get("email"), "voice")
    if not allowed:
        r = limit_response(plan_v, "voice", limit_v)
        return jsonify({"text": "", "error": r["error"], "upgrade_url": r["upgrade_url"]}), 429
    data = f.read()
    if len(data) < 1500:
        return jsonify({"text": "", "error": "That recording was too short. Tap the mic, speak, then tap again to stop."})
    mime = (f.mimetype or "audio/webm").split(";")[0]
    if not mime.startswith("audio/"):
        mime = "audio/webm"
    ext = "m4a" if ("mp4" in mime or "m4a" in mime or "aac" in mime) else ("ogg" if "ogg" in mime else
          ("wav" if "wav" in mime else ("mp3" if ("mpeg" in mime or "mp3" in mime) else "webm")))
    lang = request.form.get("language") or "Auto"
    hint = {"English": "en", "Pidgin": "en", "Hausa": "ha"}.get(lang)
    txt, errors = "", []
    if GROQ_KEY:  # Whisper accepts webm, mp4/m4a, ogg, wav, mp3 and flac
        try:
            txt = groq_transcribe(data, f"voice.{ext}", mime, hint)
        except Exception as e:
            errors.append(f"groq: {str(e)[:200]}")
            print(f"[voice] groq whisper failed: {str(e)[:300]}")
    if not txt and gclient:
        try:
            resp = gemini_generate([types.Content(role="user", parts=[
                types.Part.from_bytes(data=data, mime_type=mime),
                types.Part(text="Transcribe this audio exactly as spoken, in its original language "
                                "(English, Pidgin, Twi, Hausa or other). Return only the transcription.")])])
            txt = (resp.text or "").strip()
        except Exception as e:
            errors.append(f"gemini: {str(e)[:200]}")
            print(f"[voice] gemini failed: {str(e)[:300]}")
    txt = clean_answer(txt)
    if txt:
        add_usage(usage_key(uid_v), "voice")
        return jsonify({"text": txt})
    if not (GROQ_KEY or gclient):
        msg = "Voice typing is not set up on the server yet (it needs GROQ_API_KEY or GEMINI_API_KEY)."
    elif errors:
        msg = "Voice typing is busy right now. Please try again in a moment, or type your message."
    else:
        msg = "I could not hear any speech in that recording. Please speak closer to the mic and try again."
    return jsonify({"text": "", "error": msg})


@app.route("/media/<filename>")
def media(filename):
    p = os.path.join(GENERATED_DIR, secure_filename(filename))
    if os.path.exists(p):
        return send_file(p, mimetype="image/jpeg")
    return "not found", 404


@app.route("/download/<filename>")
def dl(filename):
    safe = secure_filename(filename)
    p = os.path.join(GENERATED_DIR, safe)
    if os.path.exists(p):
        return send_file(p, as_attachment=True, download_name=friendly_name(safe))
    return "not found", 404


@app.route("/audio/<filename>")
def au(filename):
    p = os.path.join(GENERATED_DIR, secure_filename(filename))
    if os.path.exists(p):
        return send_file(p, mimetype="audio/mpeg")
    return "not found", 404


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
